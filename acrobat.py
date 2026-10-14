import os
import sys
import time
import webbrowser
import subprocess
import tkinter as tk
from tkinter import ttk
from tkinter import filedialog
from tkinter import messagebox
from tkinter import simpledialog
from tkinter import colorchooser

try:
    import fitz
    import pypdf
    from PIL import Image
    from PIL import ImageDraw
    from PIL import ImageTk
except ImportError:
    print("Some required libraries are missing.")
    print("Install them with:")
    print("pip install pymupdf pypdf pillow reportlab python-docx pytesseract")
    sys.exit()

try:
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.lib.pagesizes import letter
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import pytesseract
    HAS_OCR = True
except ImportError:
    HAS_OCR = False


APP_NAME = "Simple Acrobat"
HOME = os.path.expanduser("~")
RECENT_FILE = os.path.join(HOME, ".simple_acrobat_recent.txt")
SIGNATURE_FILE = os.path.join(HOME, ".simple_acrobat_signature.png")
PRINT_FILE = os.path.join(HOME, ".simple_acrobat_print.pdf")
BASE_SCALE = 1.33
PAGE_GAP = 14
PAGE_MARGIN = 20
MAX_UNDO = 20
ZOOM_LEVELS = [25, 50, 75, 100, 125, 150, 200, 300, 400, 600]

STAMP_NAMES = [
    "Approved",
    "Experimental",
    "Not Approved",
    "As Is",
    "Expired",
    "Not For Public Release",
    "Confidential",
    "Final",
    "Sold",
    "Departmental",
    "For Comment",
    "Top Secret",
    "Draft",
    "For Public Release",
]

DRAG_TOOLS = [
    "select",
    "highlight",
    "underline",
    "strike",
    "squiggly",
    "rect",
    "oval",
    "line",
    "arrow",
    "pen",
    "text",
    "stamp",
    "sign",
    "edit",
    "redact",
    "link",
    "crop",
    "image",
    "field_text",
    "field_check",
]

TOOL_CURSORS = {
    "hand": "fleur",
    "select": "xterm",
    "edit": "xterm",
    "note": "plus",
    "text": "xterm",
}

TOOL_HINTS = {
    "hand": "Hand: drag to scroll, click a link to follow it",
    "select": "Select: drag over text, then press Ctrl+C to copy",
    "edit": "Edit Text: click a line of text to change it, or click empty space to add text",
    "highlight": "Highlight: drag over text",
    "underline": "Underline: drag over text",
    "strike": "Strikethrough: drag over text",
    "squiggly": "Squiggly underline: drag over text",
    "note": "Sticky Note: click where the note should go",
    "text": "Text Box: click or drag to place a text box",
    "pen": "Draw: hold the mouse button and draw freely",
    "rect": "Rectangle: drag to draw",
    "oval": "Oval: drag to draw",
    "line": "Line: drag to draw",
    "arrow": "Arrow: drag to draw",
    "stamp": "Stamp: click where the stamp should go",
    "sign": "Sign: click where the signature should go",
    "redact": "Redact: drag over the area to mark, then use Tools > Apply Redactions",
    "link": "Add Link: drag over the area that should be clickable",
    "crop": "Crop: drag over the area to keep",
    "image": "Add Image: click or drag where the image should go",
    "field_text": "Text Field: drag to create a form field",
    "field_check": "Check Box: drag to create a check box",
}


def ceil_int(value):
    number = int(value)
    if number < value:
        number = number + 1
    return number


def parse_pages(text, total):
    pages = []
    text = text.replace(" ", "")
    if text == "":
        return pages
    parts = text.split(",")
    for part in parts:
        if part == "":
            continue
        if "-" in part:
            pieces = part.split("-")
            start = pieces[0]
            end = pieces[1]
            if start == "":
                start = "1"
            if end == "":
                end = str(total)
            first = int(start)
            last = int(end)
            if first > last:
                first, last = last, first
            for number in range(first, last + 1):
                if number >= 1 and number <= total:
                    if (number - 1) not in pages:
                        pages.append(number - 1)
        else:
            number = int(part)
            if number >= 1 and number <= total:
                if (number - 1) not in pages:
                    pages.append(number - 1)
    return pages


def format_size(num):
    units = ["B", "KB", "MB", "GB"]
    index = 0
    size = float(num)
    while size >= 1024 and index < 3:
        size = size / 1024
        index = index + 1
    return f"{round(size, 1)} {units[index]}"


def rgb_to_hex(color):
    r = int(color[0] * 255)
    g = int(color[1] * 255)
    b = int(color[2] * 255)
    return f"#{r:02x}{g:02x}{b:02x}"


def int_to_rgb(value):
    r = (value // 65536) % 256
    g = (value // 256) % 256
    b = value % 256
    return (r / 255, g / 255, b / 255)


def make_rect(p0, p1):
    left = min(p0.x, p1.x)
    top = min(p0.y, p1.y)
    right = max(p0.x, p1.x)
    bottom = max(p0.y, p1.y)
    return fitz.Rect(left, top, right, bottom)


def union_rect(a, b):
    left = min(a.x0, b.x0)
    top = min(a.y0, b.y0)
    right = max(a.x1, b.x1)
    bottom = max(a.y1, b.y1)
    return fitz.Rect(left, top, right, bottom)


def pick_font(name):
    lower = name.lower()
    bold = "bold" in lower or "black" in lower or "heavy" in lower
    italic = "italic" in lower or "oblique" in lower
    if "courier" in lower or "mono" in lower or "consolas" in lower:
        family = ["cour", "cobo", "coit", "cobi"]
    elif "times" in lower or "serif" in lower and "sans" not in lower or "georgia" in lower:
        family = ["tiro", "tibo", "tiit", "tibi"]
    else:
        family = ["helv", "hebo", "heit", "hebi"]
    if bold and italic:
        return family[3]
    if bold:
        return family[1]
    if italic:
        return family[2]
    return family[0]


def only_latin(text):
    return text.encode("latin-1", "replace").decode("latin-1")


def center_window(window, width, height):
    screen_w = window.winfo_screenwidth()
    screen_h = window.winfo_screenheight()
    x = int((screen_w - width) / 2)
    y = int((screen_h - height) / 3)
    if y < 0:
        y = 0
    window.geometry(f"{width}x{height}+{x}+{y}")


def make_window(app, title, width, height):
    window = tk.Toplevel(app.root)
    window.title(title)
    center_window(window, width, height)
    window.transient(app.root)
    try:
        window.update_idletasks()
        window.grab_set()
    except Exception:
        pass
    return window


def pix_to_photo(pix):
    image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    return ImageTk.PhotoImage(image)


def read_pdf(path):
    file = open(path, "rb")
    data = file.read()
    file.close()
    return fitz.open(stream=data, filetype="pdf")


def load_recent():
    items = []
    try:
        file = open(RECENT_FILE, "r", encoding="utf-8")
        text = file.read()
        file.close()
        for line in text.split("\n"):
            line = line.strip()
            if line != "" and os.path.exists(line):
                items.append(line)
    except Exception:
        pass
    return items


def save_recent(items):
    try:
        file = open(RECENT_FILE, "w", encoding="utf-8")
        file.write("\n".join(items))
        file.close()
    except Exception:
        pass


def page_text_lines(page):
    lines = []
    text = page.get_text("text")
    for line in text.split("\n"):
        line = line.strip()
        if line != "":
            lines.append(line)
    return lines


def text_to_pdf(text, output):
    page_width, page_height = letter
    margin = 54
    columns = 82
    rows = []
    for line in text.split("\n"):
        line = only_latin(line.replace("\t", "    ").replace("\r", ""))
        while len(line) > columns:
            rows.append(line[:columns])
            line = line[columns:]
        rows.append(line)
    sheet = rl_canvas.Canvas(output, pagesize=letter)
    sheet.setFont("Courier", 10)
    y = page_height - margin
    for row in rows:
        if y < margin:
            sheet.showPage()
            sheet.setFont("Courier", 10)
            y = page_height - margin
        sheet.drawString(margin, y, row)
        y = y - 13
    sheet.save()


class ScrollFrame:
    def __init__(self, parent):
        self.holder = ttk.Frame(parent)
        self.holder.pack(fill="both", expand=True)
        self.canvas = tk.Canvas(self.holder, highlightthickness=0)
        self.bar = ttk.Scrollbar(self.holder, orient="vertical", command=self.canvas.yview)
        self.inner = ttk.Frame(self.canvas)
        self.canvas.configure(yscrollcommand=self.bar.set)
        self.bar.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True)
        self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", self.on_inner)
        self.canvas.bind("<MouseWheel>", self.on_wheel)
        self.canvas.bind("<Button-4>", self.on_wheel)
        self.canvas.bind("<Button-5>", self.on_wheel)

    def on_inner(self, event):
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))

    def on_wheel(self, event):
        if event.num == 4:
            step = -2
        elif event.num == 5:
            step = 2
        else:
            step = int(-event.delta / 120 * 2)
            if step == 0:
                if event.delta > 0:
                    step = -1
                else:
                    step = 1
        self.canvas.yview_scroll(step, "units")


class ProgressWindow:
    def __init__(self, app, title, total):
        self.app = app
        self.win = tk.Toplevel(app.root)
        self.win.title(title)
        center_window(self.win, 360, 110)
        self.win.resizable(False, False)
        self.win.transient(app.root)
        self.label = ttk.Label(self.win, text="Working...")
        self.label.pack(pady=(16, 6))
        self.bar = ttk.Progressbar(self.win, length=300, maximum=max(total, 1))
        self.bar.pack(pady=6)
        self.win.update()

    def step(self, value, text=""):
        self.bar["value"] = value
        if text != "":
            self.label.config(text=text)
        self.win.update()

    def close(self):
        self.win.destroy()


class Thumbnails:
    def __init__(self, app, parent):
        self.app = app
        self.frame = ttk.Frame(parent)
        self.canvas = tk.Canvas(self.frame, bg="#dddddd", width=145, highlightthickness=0, yscrollincrement=40)
        self.bar = ttk.Scrollbar(self.frame, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=self.on_scroll)
        self.bar.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True)
        self.positions = []
        self.photos = {}
        self.items = {}
        self.total = 1
        self.current = -1
        self.canvas.bind("<Button-1>", self.on_click)
        self.canvas.bind("<MouseWheel>", self.on_wheel)
        self.canvas.bind("<Button-4>", self.on_wheel)
        self.canvas.bind("<Button-5>", self.on_wheel)

    def on_scroll(self, first, last):
        self.bar.set(first, last)
        self.render_visible()

    def on_wheel(self, event):
        if event.num == 4:
            step = -3
        elif event.num == 5:
            step = 3
        else:
            step = int(-event.delta / 120 * 3)
            if step == 0:
                if event.delta > 0:
                    step = -1
                else:
                    step = 1
        self.canvas.yview_scroll(step, "units")

    def reload(self):
        self.canvas.delete("all")
        self.photos = {}
        self.items = {}
        self.positions = []
        self.current = -1
        doc = self.app.doc
        if doc is None:
            self.canvas.configure(scrollregion=(0, 0, 145, 1))
            return
        y = 10
        for i in range(doc.page_count):
            rect = doc[i].rect
            w = 100
            h = ceil_int(rect.height * w / rect.width)
            x = 22
            self.positions.append((x, y, w, h))
            self.canvas.create_rectangle(x, y, x + w, y + h, fill="white", outline="#999999", tags="back")
            self.canvas.create_text(x + w / 2, y + h + 12, text=str(i + 1), fill="#333333")
            y = y + h + 30
        self.total = y
        self.canvas.configure(scrollregion=(0, 0, 145, y))
        self.render_visible()

    def render_visible(self):
        if self.app.doc is None:
            return
        top = self.canvas.canvasy(0)
        bottom = top + self.canvas.winfo_height()
        for i in range(len(self.positions)):
            x, y, w, h = self.positions[i]
            if y + h >= top - 150 and y <= bottom + 150:
                if i not in self.photos:
                    self.draw_thumb(i)

    def draw_thumb(self, i):
        x, y, w, h = self.positions[i]
        try:
            page = self.app.doc[i]
            scale = w / page.rect.width
            pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            photo = pix_to_photo(pix)
            self.photos[i] = photo
            self.items[i] = self.canvas.create_image(x, y, anchor="nw", image=photo, tags="thumb")
            self.canvas.tag_raise("cur")
        except Exception:
            self.photos[i] = None

    def refresh(self, i):
        if i in self.items:
            self.canvas.delete(self.items[i])
            del self.items[i]
        if i in self.photos:
            del self.photos[i]
        self.render_visible()

    def highlight(self, i):
        if i == self.current:
            return
        if i < 0 or i >= len(self.positions):
            return
        self.current = i
        self.canvas.delete("cur")
        x, y, w, h = self.positions[i]
        self.canvas.create_rectangle(x - 3, y - 3, x + w + 3, y + h + 3, outline="#2a7de1", width=3, tags="cur")
        top = self.canvas.canvasy(0)
        bottom = top + self.canvas.winfo_height()
        if y < top or y + h > bottom:
            self.canvas.yview_moveto(max(y - 20, 0) / self.total)

    def on_click(self, event):
        if self.app.doc is None:
            return
        cy = self.canvas.canvasy(event.y)
        for i in range(len(self.positions)):
            x, y, w, h = self.positions[i]
            if cy >= y - 5 and cy <= y + h + 25:
                self.app.viewer.goto_page(i)
                return


class Viewer:
    def __init__(self, app, parent):
        self.app = app
        self.frame = ttk.Frame(parent)
        self.canvas = tk.Canvas(self.frame, bg="#7a7a7a", highlightthickness=0, yscrollincrement=40, xscrollincrement=40)
        self.vbar = ttk.Scrollbar(self.frame, orient="vertical", command=self.canvas.yview)
        self.hbar = ttk.Scrollbar(self.frame, orient="horizontal", command=self.canvas.xview)
        self.canvas.configure(yscrollcommand=self.on_yscroll, xscrollcommand=self.hbar.set)
        self.vbar.pack(side="right", fill="y")
        self.hbar.pack(side="bottom", fill="x")
        self.canvas.pack(side="left", fill="both", expand=True)
        self.percent = 100
        self.zoom = BASE_SCALE
        self.positions = []
        self.photos = {}
        self.image_ids = {}
        self.find_marks = {}
        self.find_current = None
        self.sel_marks = {}
        self.total_height = 1
        self.resize_job = None
        self.last_width = 0
        self.start = None
        self.start_page = None
        self.press_xy = (0, 0)
        self.points = []
        self.canvas.bind("<ButtonPress-1>", self.on_press)
        self.canvas.bind("<B1-Motion>", self.on_drag)
        self.canvas.bind("<ButtonRelease-1>", self.on_release)
        self.canvas.bind("<MouseWheel>", self.on_wheel)
        self.canvas.bind("<Button-4>", self.on_wheel)
        self.canvas.bind("<Button-5>", self.on_wheel)
        self.canvas.bind("<Control-MouseWheel>", self.on_ctrl_wheel)
        self.canvas.bind("<Control-Button-4>", self.on_ctrl_wheel)
        self.canvas.bind("<Control-Button-5>", self.on_ctrl_wheel)
        self.canvas.bind("<Configure>", self.on_resize)
        self.canvas.bind("<Enter>", self.on_enter)

    def on_enter(self, event):
        self.canvas.focus_set()

    def on_yscroll(self, first, last):
        self.vbar.set(first, last)
        self.render_visible()
        self.app.page_changed()

    def on_wheel(self, event):
        if self.app.doc is None:
            return
        if event.num == 4:
            step = -3
        elif event.num == 5:
            step = 3
        else:
            step = int(-event.delta / 120 * 3)
            if step == 0:
                if event.delta > 0:
                    step = -1
                else:
                    step = 1
        self.canvas.yview_scroll(step, "units")

    def on_ctrl_wheel(self, event):
        if self.app.doc is None:
            return
        zoom_in = False
        if event.num == 4:
            zoom_in = True
        elif event.num == 5:
            zoom_in = False
        elif event.delta > 0:
            zoom_in = True
        if zoom_in:
            self.app.zoom_in()
        else:
            self.app.zoom_out()

    def on_resize(self, event):
        if self.resize_job is not None:
            self.canvas.after_cancel(self.resize_job)
        self.resize_job = self.canvas.after(150, self.after_resize)

    def after_resize(self):
        self.resize_job = None
        if self.app.doc is None:
            return
        width = self.canvas.winfo_width()
        if width != self.last_width:
            self.last_width = width
            self.reload(True)
        else:
            self.render_visible()

    def set_cursor(self, tool):
        cursor = TOOL_CURSORS.get(tool, "crosshair")
        self.canvas.configure(cursor=cursor)

    def layout(self):
        self.canvas.delete("all")
        self.photos = {}
        self.image_ids = {}
        self.positions = []
        doc = self.app.doc
        if doc is None:
            self.canvas.configure(scrollregion=(0, 0, 1, 1))
            self.canvas.create_text(40, 40, anchor="nw", fill="white", font=("Arial", 16), text="Open a PDF with File > Open (Ctrl+O)", tags="welcome")
            return
        widths = []
        heights = []
        for i in range(doc.page_count):
            rect = doc[i].rect
            widths.append(ceil_int(rect.width * self.zoom))
            heights.append(ceil_int(rect.height * self.zoom))
        biggest = max(widths)
        view_width = self.canvas.winfo_width()
        total_width = biggest + 2 * PAGE_MARGIN
        if total_width < view_width:
            total_width = view_width
        y = PAGE_MARGIN
        for i in range(len(widths)):
            x = int((total_width - widths[i]) / 2)
            self.positions.append((x, y, widths[i], heights[i]))
            self.canvas.create_rectangle(x + 4, y + 4, x + widths[i] + 4, y + heights[i] + 4, fill="#555555", outline="", tags="pagebg")
            self.canvas.create_rectangle(x, y, x + widths[i], y + heights[i], fill="white", outline="#333333", tags="pagebg")
            y = y + heights[i] + PAGE_GAP
        self.total_height = y + PAGE_MARGIN - PAGE_GAP
        self.canvas.configure(scrollregion=(0, 0, total_width, self.total_height))

    def reload(self, keep=True):
        fraction = 0
        if keep:
            fraction = self.canvas.yview()[0]
        self.layout()
        self.canvas.yview_moveto(fraction)
        self.render_visible()
        self.draw_overlays()

    def render_visible(self):
        if self.app.doc is None or len(self.positions) == 0:
            return
        top = self.canvas.canvasy(0)
        height = self.canvas.winfo_height()
        bottom = top + height
        for i in range(len(self.positions)):
            x, y, w, h = self.positions[i]
            if y + h >= top - 200 and y <= bottom + 200:
                if i not in self.photos:
                    self.render_page(i)
        keys = list(self.photos.keys())
        for i in keys:
            x, y, w, h = self.positions[i]
            if y + h < top - 3 * height or y > bottom + 3 * height:
                self.drop_page(i)

    def render_page(self, i):
        x, y, w, h = self.positions[i]
        try:
            page = self.app.doc[i]
            matrix = fitz.Matrix(self.zoom, self.zoom)
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            photo = pix_to_photo(pix)
        except Exception:
            self.photos[i] = None
            return
        self.photos[i] = photo
        item = self.canvas.create_image(x, y, anchor="nw", image=photo, tags="pageimg")
        self.image_ids[i] = item
        self.canvas.tag_raise("overlay")

    def drop_page(self, i):
        if i in self.image_ids:
            self.canvas.delete(self.image_ids[i])
            del self.image_ids[i]
        if i in self.photos:
            del self.photos[i]

    def refresh_page(self, i):
        self.drop_page(i)
        self.render_visible()

    def draw_overlays(self):
        self.canvas.delete("overlay")
        for index in self.find_marks:
            if index >= len(self.positions):
                continue
            x, y, w, h = self.positions[index]
            rects = self.find_marks[index]
            for n in range(len(rects)):
                r = rects[n]
                color = "#f5a623"
                width = 2
                if self.find_current == (index, n):
                    color = "#e0301e"
                    width = 3
                self.canvas.create_rectangle(x + r.x0 * self.zoom - 1, y + r.y0 * self.zoom - 1, x + r.x1 * self.zoom + 1, y + r.y1 * self.zoom + 1, outline=color, width=width, tags="overlay")
        for index in self.sel_marks:
            if index >= len(self.positions):
                continue
            x, y, w, h = self.positions[index]
            for r in self.sel_marks[index]:
                self.canvas.create_rectangle(x + r.x0 * self.zoom, y + r.y0 * self.zoom, x + r.x1 * self.zoom, y + r.y1 * self.zoom, outline="#2a7de1", width=2, tags="overlay")
        self.canvas.tag_raise("overlay")

    def clear_marks(self):
        self.find_marks = {}
        self.find_current = None
        self.sel_marks = {}
        self.draw_overlays()

    def page_at(self, cy):
        for i in range(len(self.positions)):
            x, y, w, h = self.positions[i]
            if cy >= y and cy <= y + h:
                return i
        return None

    def current_page(self):
        if len(self.positions) == 0:
            return 0
        middle = self.canvas.canvasy(self.canvas.winfo_height() / 3)
        for i in range(len(self.positions)):
            x, y, w, h = self.positions[i]
            if middle < y + h + PAGE_GAP / 2:
                return i
        return len(self.positions) - 1

    def goto_page(self, i):
        if len(self.positions) == 0:
            return
        if i < 0:
            i = 0
        if i >= len(self.positions):
            i = len(self.positions) - 1
        y = self.positions[i][1]
        target = y - 8
        if target < 0:
            target = 0
        self.canvas.yview_moveto(target / self.total_height)
        self.render_visible()
        self.app.page_changed()

    def scroll_to_rect(self, index, rect):
        x, y, w, h = self.positions[index]
        target = y + rect.y0 * self.zoom - self.canvas.winfo_height() / 3
        if target < 0:
            target = 0
        self.canvas.yview_moveto(target / self.total_height)
        self.render_visible()
        self.app.page_changed()

    def set_percent(self, value):
        if value < 10:
            value = 10
        if value > 800:
            value = 800
        self.percent = int(value)
        self.zoom = self.percent / 100 * BASE_SCALE
        self.app.percent_var.set(f"{self.percent}%")
        if self.app.doc is None:
            return
        self.reload(True)
        self.app.page_changed()

    def fit_width(self):
        if self.app.doc is None:
            return
        page = self.app.doc[self.current_page()]
        available = self.canvas.winfo_width() - 2 * PAGE_MARGIN - 4
        percent = available / (page.rect.width * BASE_SCALE) * 100
        self.set_percent(int(percent))

    def fit_page(self):
        if self.app.doc is None:
            return
        page = self.app.doc[self.current_page()]
        available_w = self.canvas.winfo_width() - 2 * PAGE_MARGIN - 4
        available_h = self.canvas.winfo_height() - 2 * PAGE_MARGIN
        percent_w = available_w / (page.rect.width * BASE_SCALE) * 100
        percent_h = available_h / (page.rect.height * BASE_SCALE) * 100
        self.set_percent(int(min(percent_w, percent_h)))

    def on_press(self, event):
        if self.app.doc is None:
            return
        self.canvas.focus_set()
        cx = self.canvas.canvasx(event.x)
        cy = self.canvas.canvasy(event.y)
        self.start = (cx, cy)
        self.start_page = self.page_at(cy)
        self.points = [(cx, cy)]
        self.press_xy = (event.x, event.y)
        tool = self.app.tool_var.get()
        if tool == "hand":
            self.canvas.scan_mark(event.x, event.y)

    def on_drag(self, event):
        if self.start is None:
            return
        tool = self.app.tool_var.get()
        if tool == "hand":
            self.canvas.scan_dragto(event.x, event.y, gain=1)
            return
        if tool not in DRAG_TOOLS:
            return
        cx = self.canvas.canvasx(event.x)
        cy = self.canvas.canvasy(event.y)
        self.points.append((cx, cy))
        self.canvas.delete("rubber")
        sx, sy = self.start
        color = rgb_to_hex(self.app.color)
        if tool == "pen":
            flat = []
            for p in self.points:
                flat.append(p[0])
                flat.append(p[1])
            if len(flat) >= 4:
                self.canvas.create_line(flat, fill=color, width=self.app.line_width.get(), capstyle="round", tags="rubber")
        elif tool == "line" or tool == "arrow":
            arrow = None
            if tool == "arrow":
                arrow = "last"
            self.canvas.create_line(sx, sy, cx, cy, fill=color, width=self.app.line_width.get(), arrow=arrow, tags="rubber")
        elif tool == "oval":
            self.canvas.create_oval(sx, sy, cx, cy, outline=color, width=self.app.line_width.get(), tags="rubber")
        else:
            self.canvas.create_rectangle(sx, sy, cx, cy, outline="#2a7de1", dash=(4, 2), width=2, tags="rubber")

    def on_release(self, event):
        if self.start is None:
            return
        tool = self.app.tool_var.get()
        sx, sy = self.start
        cx = self.canvas.canvasx(event.x)
        cy = self.canvas.canvasy(event.y)
        index = self.start_page
        points = self.points
        self.start = None
        self.canvas.delete("rubber")
        if tool == "hand":
            moved_x = abs(event.x - self.press_xy[0])
            moved_y = abs(event.y - self.press_xy[1])
            if moved_x < 4 and moved_y < 4 and index is not None:
                x, y, w, h = self.positions[index]
                self.app.follow_link(index, (cx - x) / self.zoom, (cy - y) / self.zoom)
            return
        if index is None:
            return
        x, y, w, h = self.positions[index]
        page_w = w / self.zoom
        page_h = h / self.zoom
        clicked = True
        if abs(cx - sx) > 4 or abs(cy - sy) > 4:
            clicked = False
        start_x = (sx - x) / self.zoom
        start_y = (sy - y) / self.zoom
        end_x = (cx - x) / self.zoom
        end_y = (cy - y) / self.zoom
        end_x = min(max(end_x, 0), page_w)
        end_y = min(max(end_y, 0), page_h)
        view_points = []
        for p in points:
            px = min(max((p[0] - x) / self.zoom, 0), page_w)
            py = min(max((p[1] - y) / self.zoom, 0), page_h)
            view_points.append((px, py))
        self.app.handle_tool(index, start_x, start_y, end_x, end_y, view_points, clicked)


class MergeWindow:
    def __init__(self, app):
        self.app = app
        self.files = []
        self.win = make_window(app, "Merge Files", 560, 400)
        body = ttk.Frame(self.win)
        body.pack(fill="both", expand=True, padx=10, pady=10)
        self.listbox = tk.Listbox(body, selectmode="browse", activestyle="none")
        self.listbox.pack(side="left", fill="both", expand=True)
        side = ttk.Frame(body)
        side.pack(side="left", fill="y", padx=(10, 0))
        ttk.Button(side, text="Add Files...", command=self.add_files).pack(fill="x", pady=2)
        ttk.Button(side, text="Remove", command=self.remove_file).pack(fill="x", pady=2)
        ttk.Button(side, text="Move Up", command=self.move_up).pack(fill="x", pady=2)
        ttk.Button(side, text="Move Down", command=self.move_down).pack(fill="x", pady=2)
        ttk.Button(side, text="Clear", command=self.clear_files).pack(fill="x", pady=2)
        bottom = ttk.Frame(self.win)
        bottom.pack(fill="x", padx=10, pady=(0, 10))
        self.open_after = tk.IntVar(value=1)
        ttk.Checkbutton(bottom, text="Open the merged file", variable=self.open_after).pack(side="left")
        ttk.Button(bottom, text="Merge", command=self.merge).pack(side="right")
        ttk.Button(bottom, text="Close", command=self.win.destroy).pack(side="right", padx=6)
        if app.path is not None:
            self.files.append(app.path)
        self.fill()

    def fill(self, select=None):
        self.listbox.delete(0, "end")
        for path in self.files:
            self.listbox.insert("end", path)
        if select is not None and select >= 0 and select < len(self.files):
            self.listbox.selection_set(select)

    def selected(self):
        chosen = self.listbox.curselection()
        if len(chosen) == 0:
            return -1
        return chosen[0]

    def add_files(self):
        paths = filedialog.askopenfilenames(parent=self.win, filetypes=[("PDF files", "*.pdf")])
        for path in paths:
            self.files.append(path)
        self.fill()

    def remove_file(self):
        index = self.selected()
        if index < 0:
            return
        self.files.pop(index)
        self.fill(index)

    def move_up(self):
        index = self.selected()
        if index <= 0:
            return
        item = self.files.pop(index)
        self.files.insert(index - 1, item)
        self.fill(index - 1)

    def move_down(self):
        index = self.selected()
        if index < 0 or index >= len(self.files) - 1:
            return
        item = self.files.pop(index)
        self.files.insert(index + 1, item)
        self.fill(index + 1)

    def clear_files(self):
        self.files = []
        self.fill()

    def merge(self):
        if len(self.files) < 2:
            messagebox.showinfo("Merge Files", "Add at least two PDF files.", parent=self.win)
            return
        output = filedialog.asksaveasfilename(parent=self.win, defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")], initialfile="merged.pdf")
        if output == "":
            return
        try:
            writer = pypdf.PdfWriter()
            for path in self.files:
                reader = pypdf.PdfReader(path)
                if reader.is_encrypted:
                    password = simpledialog.askstring("Password", "Password for " + os.path.basename(path) + ":", show="*", parent=self.win)
                    if password is None:
                        return
                    reader.decrypt(password)
                for page in reader.pages:
                    writer.add_page(page)
            file = open(output, "wb")
            writer.write(file)
            file.close()
        except Exception as e:
            messagebox.showerror("Merge Files", f"Could not merge the files:\n{e}", parent=self.win)
            return
        self.win.destroy()
        messagebox.showinfo("Merge Files", "The files were merged.")
        if self.open_after.get() == 1:
            if self.app.confirm_discard():
                self.app.open_path(output)


class SplitWindow:
    def __init__(self, app):
        self.app = app
        self.win = make_window(app, "Split Document", 460, 300)
        body = ttk.Frame(self.win)
        body.pack(fill="both", expand=True, padx=14, pady=14)
        total = app.doc.page_count
        ttk.Label(body, text=f"This document has {total} pages.").pack(anchor="w", pady=(0, 10))
        self.mode = tk.StringVar(value="every")
        row = ttk.Frame(body)
        row.pack(fill="x", pady=3)
        ttk.Radiobutton(row, text="Split every", variable=self.mode, value="every").pack(side="left")
        self.every = tk.IntVar(value=1)
        ttk.Spinbox(row, from_=1, to=max(total, 1), width=5, textvariable=self.every).pack(side="left", padx=6)
        ttk.Label(row, text="pages").pack(side="left")
        row = ttk.Frame(body)
        row.pack(fill="x", pady=3)
        ttk.Radiobutton(row, text="Split by ranges", variable=self.mode, value="ranges").pack(side="left")
        self.ranges = tk.StringVar(value="1-3; 4-6")
        ttk.Entry(row, textvariable=self.ranges, width=24).pack(side="left", padx=6)
        ttk.Label(body, text="Separate ranges with semicolons, for example 1-3; 4-6; 7").pack(anchor="w", pady=(0, 6))
        ttk.Radiobutton(body, text="Split into single pages", variable=self.mode, value="each").pack(anchor="w", pady=3)
        bottom = ttk.Frame(self.win)
        bottom.pack(fill="x", padx=14, pady=(0, 14))
        ttk.Button(bottom, text="Split", command=self.split).pack(side="right")
        ttk.Button(bottom, text="Close", command=self.win.destroy).pack(side="right", padx=6)

    def split(self):
        doc = self.app.doc
        total = doc.page_count
        groups = []
        mode = self.mode.get()
        try:
            if mode == "each":
                for i in range(total):
                    groups.append([i])
            elif mode == "every":
                size = int(self.every.get())
                if size < 1:
                    size = 1
                start = 0
                while start < total:
                    group = []
                    for i in range(start, min(start + size, total)):
                        group.append(i)
                    groups.append(group)
                    start = start + size
            else:
                for part in self.ranges.get().split(";"):
                    pages = parse_pages(part, total)
                    if len(pages) > 0:
                        groups.append(pages)
        except ValueError:
            messagebox.showerror("Split Document", "Please check the numbers you typed.", parent=self.win)
            return
        if len(groups) == 0:
            messagebox.showinfo("Split Document", "Nothing to split.", parent=self.win)
            return
        folder = filedialog.askdirectory(parent=self.win, title="Choose a folder for the new files")
        if folder == "":
            return
        base = "document"
        if self.app.path is not None:
            base = os.path.splitext(os.path.basename(self.app.path))[0]
        progress = ProgressWindow(self.app, "Splitting", len(groups))
        try:
            for n in range(len(groups)):
                part = fitz.open()
                for page_index in groups[n]:
                    part.insert_pdf(doc, from_page=page_index, to_page=page_index)
                part.save(os.path.join(folder, f"{base}_part{n + 1}.pdf"))
                part.close()
                progress.step(n + 1, f"Saving part {n + 1} of {len(groups)}")
        except Exception as e:
            progress.close()
            messagebox.showerror("Split Document", f"Could not split the document:\n{e}", parent=self.win)
            return
        progress.close()
        self.win.destroy()
        messagebox.showinfo("Split Document", f"Created {len(groups)} files in:\n{folder}")


class OrganizeWindow:
    def __init__(self, app):
        self.app = app
        self.win = make_window(app, "Organize Pages", 600, 480)
        body = ttk.Frame(self.win)
        body.pack(fill="both", expand=True, padx=10, pady=10)
        self.listbox = tk.Listbox(body, selectmode="extended", activestyle="none")
        scroll = ttk.Scrollbar(body, orient="vertical", command=self.listbox.yview)
        self.listbox.configure(yscrollcommand=scroll.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        scroll.pack(side="left", fill="y")
        side = ttk.Frame(body)
        side.pack(side="left", fill="y", padx=(10, 0))
        ttk.Button(side, text="Move Up", command=self.move_up).pack(fill="x", pady=2)
        ttk.Button(side, text="Move Down", command=self.move_down).pack(fill="x", pady=2)
        ttk.Button(side, text="Rotate Right", command=self.rotate_right).pack(fill="x", pady=2)
        ttk.Button(side, text="Rotate Left", command=self.rotate_left).pack(fill="x", pady=2)
        ttk.Button(side, text="Duplicate", command=self.duplicate).pack(fill="x", pady=2)
        ttk.Button(side, text="Insert Blank", command=self.insert_blank).pack(fill="x", pady=2)
        ttk.Button(side, text="Delete", command=self.delete).pack(fill="x", pady=2)
        ttk.Button(side, text="Extract...", command=self.extract).pack(fill="x", pady=2)
        bottom = ttk.Frame(self.win)
        bottom.pack(fill="x", padx=10, pady=(0, 10))
        ttk.Label(bottom, text="Hold Ctrl or Shift to select several pages").pack(side="left")
        ttk.Button(bottom, text="Done", command=self.win.destroy).pack(side="right")
        self.fill()

    def fill(self, select=None):
        self.listbox.delete(0, "end")
        doc = self.app.doc
        for i in range(doc.page_count):
            page = doc[i]
            text = f"Page {i + 1}     {int(page.rect.width)} x {int(page.rect.height)} pt"
            if page.rotation != 0:
                text = text + f"     rotated {page.rotation} degrees"
            self.listbox.insert("end", text)
        if select is not None:
            for index in select:
                if index >= 0 and index < doc.page_count:
                    self.listbox.selection_set(index)
            if len(select) > 0:
                self.listbox.see(select[0])

    def selected(self):
        chosen = list(self.listbox.curselection())
        chosen.sort()
        return chosen

    def need_selection(self):
        chosen = self.selected()
        if len(chosen) == 0:
            messagebox.showinfo("Organize Pages", "Select one or more pages first.", parent=self.win)
        return chosen

    def finish(self, select):
        self.app.reload_all()
        self.fill(select)

    def move_up(self):
        chosen = self.need_selection()
        if len(chosen) != 1:
            if len(chosen) > 1:
                messagebox.showinfo("Organize Pages", "Select only one page to move.", parent=self.win)
            return
        index = chosen[0]
        if index == 0:
            return
        counter = self.app.push_undo()
        try:
            self.app.doc.move_page(index, index - 1)
        except Exception as e:
            self.app.rollback(counter)
            messagebox.showerror("Organize Pages", str(e), parent=self.win)
            return
        self.finish([index - 1])

    def move_down(self):
        chosen = self.need_selection()
        if len(chosen) != 1:
            if len(chosen) > 1:
                messagebox.showinfo("Organize Pages", "Select only one page to move.", parent=self.win)
            return
        index = chosen[0]
        total = self.app.doc.page_count
        if index >= total - 1:
            return
        target = index + 2
        if target >= total:
            target = -1
        counter = self.app.push_undo()
        try:
            self.app.doc.move_page(index, target)
        except Exception as e:
            self.app.rollback(counter)
            messagebox.showerror("Organize Pages", str(e), parent=self.win)
            return
        self.finish([index + 1])

    def rotate(self, amount):
        chosen = self.need_selection()
        if len(chosen) == 0:
            return
        counter = self.app.push_undo()
        try:
            for index in chosen:
                page = self.app.doc[index]
                page.set_rotation((page.rotation + amount) % 360)
        except Exception as e:
            self.app.rollback(counter)
            messagebox.showerror("Organize Pages", str(e), parent=self.win)
            return
        self.finish(chosen)

    def rotate_right(self):
        self.rotate(90)

    def rotate_left(self):
        self.rotate(270)

    def duplicate(self):
        chosen = self.need_selection()
        if len(chosen) == 0:
            return
        counter = self.app.push_undo()
        try:
            for index in chosen[::-1]:
                self.app.doc.copy_page(index, index + 1)
        except Exception as e:
            self.app.rollback(counter)
            messagebox.showerror("Organize Pages", str(e), parent=self.win)
            return
        self.finish(chosen)

    def insert_blank(self):
        chosen = self.need_selection()
        if len(chosen) == 0:
            return
        index = chosen[-1]
        rect = self.app.doc[index].rect
        counter = self.app.push_undo()
        try:
            self.app.doc.new_page(pno=index + 1, width=rect.width, height=rect.height)
        except Exception as e:
            self.app.rollback(counter)
            messagebox.showerror("Organize Pages", str(e), parent=self.win)
            return
        self.finish([index + 1])

    def delete(self):
        chosen = self.need_selection()
        if len(chosen) == 0:
            return
        if len(chosen) >= self.app.doc.page_count:
            messagebox.showinfo("Organize Pages", "A document needs at least one page.", parent=self.win)
            return
        if not messagebox.askyesno("Delete Pages", f"Delete {len(chosen)} page(s)?", parent=self.win):
            return
        counter = self.app.push_undo()
        try:
            for index in chosen[::-1]:
                self.app.doc.delete_page(index)
        except Exception as e:
            self.app.rollback(counter)
            messagebox.showerror("Organize Pages", str(e), parent=self.win)
            return
        self.finish([])

    def extract(self):
        chosen = self.need_selection()
        if len(chosen) == 0:
            return
        output = filedialog.asksaveasfilename(parent=self.win, defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")], initialfile="extracted.pdf")
        if output == "":
            return
        try:
            new_doc = fitz.open()
            for index in chosen:
                new_doc.insert_pdf(self.app.doc, from_page=index, to_page=index)
            new_doc.save(output)
            new_doc.close()
        except Exception as e:
            messagebox.showerror("Organize Pages", str(e), parent=self.win)
            return
        messagebox.showinfo("Organize Pages", "The pages were saved to a new file.", parent=self.win)


class FormWindow:
    def __init__(self, app):
        self.app = app
        self.fields = []
        doc = app.doc
        for i in range(doc.page_count):
            page = doc[i]
            for widget in page.widgets():
                kind = widget.field_type_string
                if kind == "Button" or kind == "Signature":
                    continue
                field = {}
                field["page"] = i
                field["xref"] = widget.xref
                field["name"] = widget.field_name
                field["kind"] = kind
                field["value"] = widget.field_value
                field["choices"] = widget.choice_values
                field["var"] = None
                self.fields.append(field)
        self.win = make_window(app, "Fill Form", 520, 560)
        bottom = ttk.Frame(self.win)
        bottom.pack(side="bottom", fill="x", padx=10, pady=10)
        ttk.Button(bottom, text="Apply", command=self.apply).pack(side="right")
        ttk.Button(bottom, text="Cancel", command=self.win.destroy).pack(side="right", padx=6)
        scroller = ScrollFrame(self.win)
        inner = scroller.inner
        row = 0
        for field in self.fields:
            label = f"{field['name']}  (page {field['page'] + 1})"
            ttk.Label(inner, text=label).grid(row=row, column=0, sticky="w", padx=10, pady=5)
            kind = field["kind"]
            value = field["value"]
            if kind == "CheckBox" or kind == "RadioButton":
                text = str(value)
                checked = text not in ("False", "Off", "", "None", "0")
                var = tk.BooleanVar(value=checked)
                ttk.Checkbutton(inner, variable=var).grid(row=row, column=1, sticky="w", padx=10)
            elif kind == "ComboBox" or kind == "ListBox":
                text = ""
                if value is not None:
                    text = str(value)
                var = tk.StringVar(value=text)
                choices = field["choices"]
                if choices is None:
                    choices = []
                ttk.Combobox(inner, textvariable=var, values=choices, width=28).grid(row=row, column=1, sticky="w", padx=10)
            else:
                text = ""
                if value is not None:
                    text = str(value)
                var = tk.StringVar(value=text)
                ttk.Entry(inner, textvariable=var, width=32).grid(row=row, column=1, sticky="w", padx=10)
            field["var"] = var
            row = row + 1

    def apply(self):
        counter = self.app.push_undo()
        try:
            for field in self.fields:
                page = self.app.doc[field["page"]]
                for widget in page.widgets():
                    if widget.xref == field["xref"]:
                        kind = field["kind"]
                        if kind == "CheckBox" or kind == "RadioButton":
                            widget.field_value = bool(field["var"].get())
                        else:
                            widget.field_value = field["var"].get()
                        widget.update()
        except Exception as e:
            self.app.rollback(counter)
            messagebox.showerror("Fill Form", f"Could not fill the form:\n{e}", parent=self.win)
            return
        self.win.destroy()
        self.app.reload_all()
        self.app.set_status("Form updated")


def unrotated_rect(page):
    return page.rect * page.derotation_matrix


class WatermarkWindow:
    def __init__(self, app):
        self.app = app
        self.color = (0.6, 0.6, 0.6)
        self.win = make_window(app, "Add Watermark", 460, 360)
        body = ttk.Frame(self.win)
        body.pack(fill="both", expand=True, padx=16, pady=16)
        self.text = tk.StringVar(value="CONFIDENTIAL")
        self.size = tk.IntVar(value=60)
        self.angle = tk.IntVar(value=45)
        self.opacity = tk.IntVar(value=30)
        self.pages = tk.StringVar(value="")
        ttk.Label(body, text="Watermark text").grid(row=0, column=0, sticky="w", pady=6)
        ttk.Entry(body, textvariable=self.text, width=30).grid(row=0, column=1, sticky="w", pady=6)
        ttk.Label(body, text="Font size").grid(row=1, column=0, sticky="w", pady=6)
        ttk.Spinbox(body, from_=8, to=300, width=8, textvariable=self.size).grid(row=1, column=1, sticky="w", pady=6)
        ttk.Label(body, text="Angle (degrees)").grid(row=2, column=0, sticky="w", pady=6)
        ttk.Spinbox(body, from_=-90, to=90, width=8, textvariable=self.angle).grid(row=2, column=1, sticky="w", pady=6)
        ttk.Label(body, text="Opacity (percent)").grid(row=3, column=0, sticky="w", pady=6)
        ttk.Spinbox(body, from_=5, to=100, width=8, textvariable=self.opacity).grid(row=3, column=1, sticky="w", pady=6)
        ttk.Label(body, text="Color").grid(row=4, column=0, sticky="w", pady=6)
        color_row = ttk.Frame(body)
        color_row.grid(row=4, column=1, sticky="w", pady=6)
        self.swatch = tk.Label(color_row, text="      ", bg=rgb_to_hex(self.color), relief="sunken")
        self.swatch.pack(side="left")
        ttk.Button(color_row, text="Choose...", command=self.choose_color).pack(side="left", padx=8)
        ttk.Label(body, text="Pages (blank for all)").grid(row=5, column=0, sticky="w", pady=6)
        ttk.Entry(body, textvariable=self.pages, width=18).grid(row=5, column=1, sticky="w", pady=6)
        bottom = ttk.Frame(self.win)
        bottom.pack(fill="x", padx=16, pady=(0, 16))
        ttk.Button(bottom, text="Add Watermark", command=self.apply).pack(side="right")
        ttk.Button(bottom, text="Cancel", command=self.win.destroy).pack(side="right", padx=6)

    def choose_color(self):
        result = colorchooser.askcolor(parent=self.win, initialcolor=rgb_to_hex(self.color))
        if result[0] is None:
            return
        self.color = (result[0][0] / 255, result[0][1] / 255, result[0][2] / 255)
        self.swatch.config(bg=rgb_to_hex(self.color))

    def apply(self):
        doc = self.app.doc
        total = doc.page_count
        try:
            pages = parse_pages(self.pages.get(), total)
            size = int(self.size.get())
            angle = int(self.angle.get())
            opacity = int(self.opacity.get()) / 100
        except ValueError:
            messagebox.showerror("Add Watermark", "Please check the numbers you typed.", parent=self.win)
            return
        if len(pages) == 0:
            pages = list(range(total))
        text = only_latin(self.text.get())
        if text == "":
            messagebox.showinfo("Add Watermark", "Type the watermark text first.", parent=self.win)
            return
        counter = self.app.push_undo()
        try:
            length = fitz.get_text_length(text, fontname="helv", fontsize=size)
            for index in pages:
                page = doc[index]
                rect = unrotated_rect(page)
                center = fitz.Point(rect.width / 2, rect.height / 2)
                origin = fitz.Point(center.x - length / 2, center.y + size / 3)
                page.insert_text(origin, text, fontsize=size, fontname="helv", color=self.color, fill_opacity=opacity, morph=(center, fitz.Matrix(-angle)), overlay=True)
        except Exception as e:
            self.app.rollback(counter)
            messagebox.showerror("Add Watermark", f"Could not add the watermark:\n{e}", parent=self.win)
            return
        self.win.destroy()
        self.app.reload_all()
        self.app.set_status("Watermark added")


class HeaderFooterWindow:
    def __init__(self, app):
        self.app = app
        self.win = make_window(app, "Header and Footer", 600, 400)
        body = ttk.Frame(self.win)
        body.pack(fill="both", expand=True, padx=16, pady=16)
        self.entries = {}
        names = ["Left", "Center", "Right"]
        for column in range(3):
            ttk.Label(body, text=names[column]).grid(row=0, column=column + 1, sticky="w", padx=4)
        ttk.Label(body, text="Header").grid(row=1, column=0, sticky="w", pady=6)
        ttk.Label(body, text="Footer").grid(row=2, column=0, sticky="w", pady=6)
        for row_name, row in (("header", 1), ("footer", 2)):
            for column in range(3):
                var = tk.StringVar(value="")
                if row_name == "footer" and column == 1:
                    var.set("Page {n} of {total}")
                ttk.Entry(body, textvariable=var, width=20).grid(row=row, column=column + 1, padx=4, pady=6)
                self.entries[f"{row_name}{column}"] = var
        ttk.Label(body, text="You can use {n} for the page number, {total} for the page count,\n{date} for today's date and {file} for the file name.").grid(row=3, column=0, columnspan=4, sticky="w", pady=8)
        self.size = tk.IntVar(value=10)
        self.margin = tk.IntVar(value=30)
        self.start = tk.IntVar(value=1)
        self.pages = tk.StringVar(value="")
        ttk.Label(body, text="Font size").grid(row=4, column=0, sticky="w", pady=4)
        ttk.Spinbox(body, from_=6, to=48, width=6, textvariable=self.size).grid(row=4, column=1, sticky="w", padx=4)
        ttk.Label(body, text="Margin").grid(row=5, column=0, sticky="w", pady=4)
        ttk.Spinbox(body, from_=10, to=150, width=6, textvariable=self.margin).grid(row=5, column=1, sticky="w", padx=4)
        ttk.Label(body, text="First page number").grid(row=6, column=0, sticky="w", pady=4)
        ttk.Spinbox(body, from_=0, to=9999, width=6, textvariable=self.start).grid(row=6, column=1, sticky="w", padx=4)
        ttk.Label(body, text="Pages (blank for all)").grid(row=7, column=0, sticky="w", pady=4)
        ttk.Entry(body, textvariable=self.pages, width=14).grid(row=7, column=1, sticky="w", padx=4)
        bottom = ttk.Frame(self.win)
        bottom.pack(fill="x", padx=16, pady=(0, 16))
        ttk.Button(bottom, text="Apply", command=self.apply).pack(side="right")
        ttk.Button(bottom, text="Cancel", command=self.win.destroy).pack(side="right", padx=6)

    def apply(self):
        doc = self.app.doc
        total = doc.page_count
        try:
            pages = parse_pages(self.pages.get(), total)
            size = int(self.size.get())
            margin = int(self.margin.get())
            start = int(self.start.get())
        except ValueError:
            messagebox.showerror("Header and Footer", "Please check the numbers you typed.", parent=self.win)
            return
        if len(pages) == 0:
            pages = list(range(total))
        file_name = "document"
        if self.app.path is not None:
            file_name = os.path.basename(self.app.path)
        today = time.strftime("%Y-%m-%d")
        counter = self.app.push_undo()
        try:
            for index in pages:
                page = doc[index]
                rect = unrotated_rect(page)
                for row_name in ("header", "footer"):
                    for column in range(3):
                        text = self.entries[f"{row_name}{column}"].get()
                        if text == "":
                            continue
                        text = text.replace("{n}", str(index + start))
                        text = text.replace("{total}", str(total))
                        text = text.replace("{date}", today)
                        text = text.replace("{file}", file_name)
                        text = only_latin(text)
                        length = fitz.get_text_length(text, fontname="helv", fontsize=size)
                        if column == 0:
                            x = margin
                        elif column == 1:
                            x = (rect.width - length) / 2
                        else:
                            x = rect.width - margin - length
                        if row_name == "header":
                            y = margin + size
                        else:
                            y = rect.height - margin
                        page.insert_text(fitz.Point(x, y), text, fontsize=size, fontname="helv", color=(0, 0, 0))
        except Exception as e:
            self.app.rollback(counter)
            messagebox.showerror("Header and Footer", f"Could not add the text:\n{e}", parent=self.win)
            return
        self.win.destroy()
        self.app.reload_all()
        self.app.set_status("Header and footer added")


class PropertiesWindow:
    def __init__(self, app):
        self.app = app
        doc = app.doc
        self.win = make_window(app, "Document Properties", 520, 520)
        body = ttk.Frame(self.win)
        body.pack(fill="both", expand=True, padx=16, pady=16)
        meta = doc.metadata
        if meta is None:
            meta = {}
        self.keys = ["title", "author", "subject", "keywords", "creator", "producer"]
        self.vars = {}
        row = 0
        for key in self.keys:
            ttk.Label(body, text=key.capitalize()).grid(row=row, column=0, sticky="w", pady=4)
            value = meta.get(key, "")
            if value is None:
                value = ""
            var = tk.StringVar(value=value)
            ttk.Entry(body, textvariable=var, width=44).grid(row=row, column=1, sticky="w", pady=4, padx=8)
            self.vars[key] = var
            row = row + 1
        ttk.Separator(body, orient="horizontal").grid(row=row, column=0, columnspan=2, sticky="ew", pady=10)
        row = row + 1
        file_name = "Not saved yet"
        file_size = "Unknown"
        if app.path is not None:
            file_name = os.path.basename(app.path)
            if os.path.exists(app.path):
                file_size = format_size(os.path.getsize(app.path))
        page_rect = doc[0].rect
        inches_w = round(page_rect.width / 72, 2)
        inches_h = round(page_rect.height / 72, 2)
        form_fields = 0
        for i in range(doc.page_count):
            for widget in doc[i].widgets():
                form_fields = form_fields + 1
        details = [
            ("File name", file_name),
            ("File size", file_size),
            ("Pages", str(doc.page_count)),
            ("Page size", f"{inches_w} x {inches_h} in"),
            ("PDF version", str(meta.get("format", "Unknown"))),
            ("Encrypted", str(doc.is_encrypted)),
            ("Form fields", str(form_fields)),
        ]
        for name, value in details:
            ttk.Label(body, text=name).grid(row=row, column=0, sticky="w", pady=2)
            ttk.Label(body, text=value).grid(row=row, column=1, sticky="w", pady=2, padx=8)
            row = row + 1
        bottom = ttk.Frame(self.win)
        bottom.pack(fill="x", padx=16, pady=(0, 16))
        ttk.Button(bottom, text="Save Properties", command=self.save).pack(side="right")
        ttk.Button(bottom, text="Close", command=self.win.destroy).pack(side="right", padx=6)

    def save(self):
        counter = self.app.push_undo()
        try:
            meta = self.app.doc.metadata
            if meta is None:
                meta = {}
            for key in self.keys:
                meta[key] = self.vars[key].get()
            self.app.doc.set_metadata(meta)
        except Exception as e:
            self.app.rollback(counter)
            messagebox.showerror("Properties", str(e), parent=self.win)
            return
        self.win.destroy()
        self.app.set_status("Properties updated. Save the file to keep them.")


class ProtectWindow:
    def __init__(self, app):
        self.app = app
        self.win = make_window(app, "Protect with Password", 460, 380)
        body = ttk.Frame(self.win)
        body.pack(fill="both", expand=True, padx=16, pady=16)
        self.user = tk.StringVar()
        self.again = tk.StringVar()
        self.owner = tk.StringVar()
        self.can_print = tk.IntVar(value=1)
        self.can_copy = tk.IntVar(value=1)
        self.can_edit = tk.IntVar(value=0)
        self.can_comment = tk.IntVar(value=1)
        ttk.Label(body, text="Password to open").grid(row=0, column=0, sticky="w", pady=6)
        ttk.Entry(body, textvariable=self.user, show="*", width=26).grid(row=0, column=1, pady=6)
        ttk.Label(body, text="Type it again").grid(row=1, column=0, sticky="w", pady=6)
        ttk.Entry(body, textvariable=self.again, show="*", width=26).grid(row=1, column=1, pady=6)
        ttk.Label(body, text="Password to change\nthe permissions").grid(row=2, column=0, sticky="w", pady=6)
        ttk.Entry(body, textvariable=self.owner, show="*", width=26).grid(row=2, column=1, pady=6)
        ttk.Label(body, text="People who open the file may:").grid(row=3, column=0, columnspan=2, sticky="w", pady=(12, 2))
        ttk.Checkbutton(body, text="Print", variable=self.can_print).grid(row=4, column=0, sticky="w")
        ttk.Checkbutton(body, text="Copy text and images", variable=self.can_copy).grid(row=5, column=0, columnspan=2, sticky="w")
        ttk.Checkbutton(body, text="Edit the document", variable=self.can_edit).grid(row=6, column=0, columnspan=2, sticky="w")
        ttk.Checkbutton(body, text="Add comments", variable=self.can_comment).grid(row=7, column=0, columnspan=2, sticky="w")
        bottom = ttk.Frame(self.win)
        bottom.pack(fill="x", padx=16, pady=(0, 16))
        ttk.Button(bottom, text="Save Protected Copy", command=self.save).pack(side="right")
        ttk.Button(bottom, text="Cancel", command=self.win.destroy).pack(side="right", padx=6)

    def save(self):
        user = self.user.get()
        owner = self.owner.get()
        if user == "" and owner == "":
            messagebox.showinfo("Protect", "Type at least one password.", parent=self.win)
            return
        if user != self.again.get():
            messagebox.showerror("Protect", "The two passwords do not match.", parent=self.win)
            return
        if owner == "":
            owner = user + "-owner"
        permissions = fitz.PDF_PERM_ACCESSIBILITY
        if self.can_print.get() == 1:
            permissions = permissions + fitz.PDF_PERM_PRINT + fitz.PDF_PERM_PRINT_HQ
        if self.can_copy.get() == 1:
            permissions = permissions + fitz.PDF_PERM_COPY
        if self.can_edit.get() == 1:
            permissions = permissions + fitz.PDF_PERM_MODIFY + fitz.PDF_PERM_ASSEMBLE + fitz.PDF_PERM_FORM
        if self.can_comment.get() == 1:
            permissions = permissions + fitz.PDF_PERM_ANNOTATE
        output = filedialog.asksaveasfilename(parent=self.win, defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")], initialfile="protected.pdf")
        if output == "":
            return
        try:
            data = self.app.doc.tobytes()
            copy = fitz.open(stream=data, filetype="pdf")
            if copy.needs_pass:
                copy.authenticate(self.app.password)
            copy.save(output, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=owner, user_pw=user, permissions=permissions)
            copy.close()
        except Exception as e:
            messagebox.showerror("Protect", f"Could not protect the file:\n{e}", parent=self.win)
            return
        self.win.destroy()
        messagebox.showinfo("Protect", "A password protected copy was saved.\nYour open document was not changed.")


class CompressWindow:
    def __init__(self, app):
        self.app = app
        self.win = make_window(app, "Compress PDF", 420, 280)
        body = ttk.Frame(self.win)
        body.pack(fill="both", expand=True, padx=16, pady=16)
        self.level = tk.StringVar(value="medium")
        ttk.Label(body, text="Choose how small the file should become:").pack(anchor="w", pady=(0, 10))
        ttk.Radiobutton(body, text="Low - keeps the best quality", variable=self.level, value="low").pack(anchor="w", pady=3)
        ttk.Radiobutton(body, text="Medium - a good balance", variable=self.level, value="medium").pack(anchor="w", pady=3)
        ttk.Radiobutton(body, text="High - smallest file, images look softer", variable=self.level, value="high").pack(anchor="w", pady=3)
        bottom = ttk.Frame(self.win)
        bottom.pack(fill="x", padx=16, pady=(0, 16))
        ttk.Button(bottom, text="Compress", command=self.compress).pack(side="right")
        ttk.Button(bottom, text="Cancel", command=self.win.destroy).pack(side="right", padx=6)

    def compress(self):
        output = filedialog.asksaveasfilename(parent=self.win, defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")], initialfile="compressed.pdf")
        if output == "":
            return
        level = self.level.get()
        try:
            data = self.app.doc.tobytes()
            work = fitz.open(stream=data, filetype="pdf")
            if work.needs_pass:
                work.authenticate(self.app.password)
            if level == "medium":
                try:
                    work.rewrite_images(dpi_threshold=200, dpi_target=150, quality=75)
                except Exception:
                    pass
            if level == "high":
                try:
                    work.rewrite_images(dpi_threshold=120, dpi_target=96, quality=50)
                except Exception:
                    pass
                try:
                    work.subset_fonts()
                except Exception:
                    pass
            if level == "low":
                work.save(output, garbage=3, deflate=True)
            else:
                work.save(output, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
            work.close()
        except Exception as e:
            messagebox.showerror("Compress", f"Could not compress the file:\n{e}", parent=self.win)
            return
        before = len(data)
        after = os.path.getsize(output)
        self.win.destroy()
        messagebox.showinfo("Compress", f"Before: {format_size(before)}\nAfter: {format_size(after)}")


class CompareWindow:
    def __init__(self, app, other_path):
        self.app = app
        self.other_path = other_path
        self.win = make_window(app, "Compare Files", 780, 560)
        self.summary = ttk.Label(self.win, text="Comparing...")
        self.summary.pack(anchor="w", padx=12, pady=(12, 4))
        holder = ttk.Frame(self.win)
        holder.pack(fill="both", expand=True, padx=12, pady=4)
        self.text = tk.Text(holder, wrap="word", font=("Courier", 10))
        bar = ttk.Scrollbar(holder, orient="vertical", command=self.text.yview)
        self.text.configure(yscrollcommand=bar.set)
        bar.pack(side="right", fill="y")
        self.text.pack(side="left", fill="both", expand=True)
        self.text.tag_configure("removed", foreground="#b00020")
        self.text.tag_configure("added", foreground="#0a7a2f")
        self.text.tag_configure("head", font=("Courier", 10, "bold"))
        legend = ttk.Label(self.win, text="Lines starting with - are only in the open file. Lines starting with + are only in the other file.")
        legend.pack(anchor="w", padx=12, pady=4)
        ttk.Button(self.win, text="Close", command=self.win.destroy).pack(anchor="e", padx=12, pady=(0, 12))
        self.run()

    def add(self, text, tag):
        self.text.insert("end", text + "\n", tag)

    def run(self):
        try:
            other = read_pdf(self.other_path)
            if other.needs_pass:
                password = simpledialog.askstring("Password", "Password for the other file:", show="*", parent=self.win)
                if password is None or other.authenticate(password) == 0:
                    self.summary.config(text="Could not open the other file.")
                    return
        except Exception as e:
            self.summary.config(text=f"Could not open the other file: {e}")
            return
        first = self.app.doc
        name_a = "open file"
        if self.app.path is not None:
            name_a = os.path.basename(self.app.path)
        name_b = os.path.basename(self.other_path)
        differences = 0
        pages = max(first.page_count, other.page_count)
        if first.page_count != other.page_count:
            self.add(f"Page count: {name_a} has {first.page_count}, {name_b} has {other.page_count}", "head")
            differences = differences + 1
        for i in range(pages):
            if i >= first.page_count:
                self.add(f"Page {i + 1} only exists in {name_b}", "added")
                differences = differences + 1
                continue
            if i >= other.page_count:
                self.add(f"Page {i + 1} only exists in {name_a}", "removed")
                differences = differences + 1
                continue
            lines_a = page_text_lines(first[i])
            lines_b = page_text_lines(other[i])
            changes = []
            for line in lines_a:
                if line not in lines_b:
                    changes.append(("- " + line, "removed"))
            for line in lines_b:
                if line not in lines_a:
                    changes.append(("+ " + line, "added"))
            if len(changes) > 0:
                self.add(f"Page {i + 1}", "head")
                for change in changes:
                    self.add(change[0], change[1])
                self.add("", "head")
                differences = differences + len(changes)
        other.close()
        if differences == 0:
            self.summary.config(text="The text of both files is identical.")
            self.add("No differences found.", "head")
        else:
            self.summary.config(text=f"Found {differences} difference(s).")
        self.text.config(state="disabled")


class BookmarkWindow:
    def __init__(self, app):
        self.app = app
        self.toc = []
        for item in app.doc.get_toc():
            self.toc.append([item[0], item[1], item[2]])
        self.win = make_window(app, "Bookmarks", 600, 460)
        body = ttk.Frame(self.win)
        body.pack(fill="both", expand=True, padx=10, pady=10)
        self.listbox = tk.Listbox(body, selectmode="browse", activestyle="none")
        self.listbox.pack(side="left", fill="both", expand=True)
        side = ttk.Frame(body)
        side.pack(side="left", fill="y", padx=(10, 0))
        ttk.Button(side, text="Add Here", command=self.add).pack(fill="x", pady=2)
        ttk.Button(side, text="Rename", command=self.rename).pack(fill="x", pady=2)
        ttk.Button(side, text="Change Page", command=self.change_page).pack(fill="x", pady=2)
        ttk.Button(side, text="Indent", command=self.indent).pack(fill="x", pady=2)
        ttk.Button(side, text="Outdent", command=self.outdent).pack(fill="x", pady=2)
        ttk.Button(side, text="Move Up", command=self.move_up).pack(fill="x", pady=2)
        ttk.Button(side, text="Move Down", command=self.move_down).pack(fill="x", pady=2)
        ttk.Button(side, text="Delete", command=self.delete).pack(fill="x", pady=2)
        bottom = ttk.Frame(self.win)
        bottom.pack(fill="x", padx=10, pady=(0, 10))
        ttk.Button(bottom, text="Save Bookmarks", command=self.save).pack(side="right")
        ttk.Button(bottom, text="Cancel", command=self.win.destroy).pack(side="right", padx=6)
        self.fill()

    def fill(self, select=None):
        self.listbox.delete(0, "end")
        for item in self.toc:
            indent = "      " * (item[0] - 1)
            self.listbox.insert("end", f"{indent}{item[1]}   (page {item[2]})")
        if select is not None and select >= 0 and select < len(self.toc):
            self.listbox.selection_set(select)
            self.listbox.see(select)

    def selected(self):
        chosen = self.listbox.curselection()
        if len(chosen) == 0:
            return -1
        return chosen[0]

    def add(self):
        title = simpledialog.askstring("Add Bookmark", "Bookmark name:", parent=self.win)
        if title is None or title.strip() == "":
            return
        page_number = self.app.viewer.current_page() + 1
        index = self.selected()
        level = 1
        if index >= 0:
            level = self.toc[index][0]
            self.toc.insert(index + 1, [level, title.strip(), page_number])
            self.fill(index + 1)
        else:
            self.toc.append([level, title.strip(), page_number])
            self.fill(len(self.toc) - 1)

    def rename(self):
        index = self.selected()
        if index < 0:
            return
        title = simpledialog.askstring("Rename Bookmark", "New name:", initialvalue=self.toc[index][1], parent=self.win)
        if title is None or title.strip() == "":
            return
        self.toc[index][1] = title.strip()
        self.fill(index)

    def change_page(self):
        index = self.selected()
        if index < 0:
            return
        total = self.app.doc.page_count
        number = simpledialog.askinteger("Change Page", f"Page number (1 to {total}):", initialvalue=self.toc[index][2], minvalue=1, maxvalue=total, parent=self.win)
        if number is None:
            return
        self.toc[index][2] = number
        self.fill(index)

    def indent(self):
        index = self.selected()
        if index <= 0:
            return
        if self.toc[index][0] <= self.toc[index - 1][0]:
            self.toc[index][0] = self.toc[index][0] + 1
        self.fill(index)

    def outdent(self):
        index = self.selected()
        if index < 0:
            return
        if self.toc[index][0] > 1:
            self.toc[index][0] = self.toc[index][0] - 1
        self.fill(index)

    def move_up(self):
        index = self.selected()
        if index <= 0:
            return
        item = self.toc.pop(index)
        self.toc.insert(index - 1, item)
        self.fill(index - 1)

    def move_down(self):
        index = self.selected()
        if index < 0 or index >= len(self.toc) - 1:
            return
        item = self.toc.pop(index)
        self.toc.insert(index + 1, item)
        self.fill(index + 1)

    def delete(self):
        index = self.selected()
        if index < 0:
            return
        self.toc.pop(index)
        self.fill(min(index, len(self.toc) - 1))

    def save(self):
        previous = 0
        for item in self.toc:
            if item[0] < 1:
                item[0] = 1
            if item[0] > previous + 1:
                item[0] = previous + 1
            previous = item[0]
        counter = self.app.push_undo()
        try:
            self.app.doc.set_toc(self.toc)
        except Exception as e:
            self.app.rollback(counter)
            messagebox.showerror("Bookmarks", str(e), parent=self.win)
            return
        self.win.destroy()
        self.app.update_panels()
        self.app.set_status("Bookmarks updated")


class SignWindow:
    def __init__(self, app):
        self.app = app
        self.strokes = []
        self.current = []
        self.image_path = tk.StringVar(value="")
        self.typed = tk.StringVar(value="")
        self.win = make_window(app, "Sign", 520, 420)
        self.tabs = ttk.Notebook(self.win)
        self.tabs.pack(fill="both", expand=True, padx=10, pady=10)
        draw_tab = ttk.Frame(self.tabs)
        type_tab = ttk.Frame(self.tabs)
        image_tab = ttk.Frame(self.tabs)
        self.tabs.add(draw_tab, text="Draw")
        self.tabs.add(type_tab, text="Type")
        self.tabs.add(image_tab, text="Image")
        ttk.Label(draw_tab, text="Draw your signature below").pack(anchor="w", padx=8, pady=6)
        self.pad = tk.Canvas(draw_tab, width=440, height=170, bg="white", highlightthickness=1, highlightbackground="#888888")
        self.pad.pack(padx=8)
        self.pad.bind("<ButtonPress-1>", self.pad_press)
        self.pad.bind("<B1-Motion>", self.pad_drag)
        self.pad.bind("<ButtonRelease-1>", self.pad_release)
        row = ttk.Frame(draw_tab)
        row.pack(fill="x", padx=8, pady=8)
        ttk.Button(row, text="Clear", command=self.clear_pad).pack(side="left")
        if os.path.exists(SIGNATURE_FILE):
            ttk.Button(row, text="Use Saved Signature", command=self.use_saved).pack(side="left", padx=8)
        ttk.Label(type_tab, text="Type your name").pack(anchor="w", padx=8, pady=(14, 6))
        ttk.Entry(type_tab, textvariable=self.typed, width=36, font=("Times", 18, "italic")).pack(padx=8, pady=6)
        ttk.Label(image_tab, text="Choose a picture of your signature").pack(anchor="w", padx=8, pady=(14, 6))
        ttk.Entry(image_tab, textvariable=self.image_path, width=48).pack(padx=8, pady=6)
        ttk.Button(image_tab, text="Browse...", command=self.browse).pack(anchor="w", padx=8)
        bottom = ttk.Frame(self.win)
        bottom.pack(fill="x", padx=10, pady=(0, 10))
        ttk.Button(bottom, text="Use Signature", command=self.use).pack(side="right")
        ttk.Button(bottom, text="Cancel", command=self.win.destroy).pack(side="right", padx=6)

    def pad_press(self, event):
        self.current = [(event.x, event.y)]

    def pad_drag(self, event):
        if len(self.current) == 0:
            return
        last = self.current[-1]
        self.pad.create_line(last[0], last[1], event.x, event.y, width=3, fill="#00003c", capstyle="round")
        self.current.append((event.x, event.y))

    def pad_release(self, event):
        if len(self.current) > 0:
            self.strokes.append(self.current)
        self.current = []

    def clear_pad(self):
        self.pad.delete("all")
        self.strokes = []
        self.current = []

    def browse(self):
        path = filedialog.askopenfilename(parent=self.win, filetypes=[("Images", "*.png *.jpg *.jpeg *.bmp *.gif"), ("All files", "*.*")])
        if path != "":
            self.image_path.set(path)

    def use_saved(self):
        self.app.sign_mode = "image"
        self.app.sign_path = SIGNATURE_FILE
        self.win.destroy()
        self.app.set_tool("sign")

    def use(self):
        tab = self.tabs.index(self.tabs.select())
        if tab == 0:
            if len(self.strokes) == 0:
                messagebox.showinfo("Sign", "Draw your signature first.", parent=self.win)
                return
            image = Image.new("RGBA", (440, 170), (0, 0, 0, 0))
            pen = ImageDraw.Draw(image)
            for stroke in self.strokes:
                if len(stroke) == 1:
                    x, y = stroke[0]
                    pen.ellipse((x - 1, y - 1, x + 1, y + 1), fill=(0, 0, 60, 255))
                else:
                    pen.line(stroke, fill=(0, 0, 60, 255), width=3, joint="curve")
            box = image.getbbox()
            if box is not None:
                image = image.crop(box)
            try:
                image.save(SIGNATURE_FILE)
            except Exception as e:
                messagebox.showerror("Sign", str(e), parent=self.win)
                return
            self.app.sign_mode = "image"
            self.app.sign_path = SIGNATURE_FILE
        elif tab == 1:
            text = self.typed.get().strip()
            if text == "":
                messagebox.showinfo("Sign", "Type your name first.", parent=self.win)
                return
            self.app.sign_mode = "text"
            self.app.sign_text = text
        else:
            path = self.image_path.get().strip()
            if path == "" or not os.path.exists(path):
                messagebox.showinfo("Sign", "Choose an image file first.", parent=self.win)
                return
            self.app.sign_mode = "image"
            self.app.sign_path = path
        self.win.destroy()
        self.app.set_tool("sign")


class StampWindow:
    def __init__(self, app):
        self.app = app
        self.win = make_window(app, "Stamp", 400, 460)
        body = ttk.Frame(self.win)
        body.pack(fill="both", expand=True, padx=14, pady=14)
        ttk.Label(body, text="Choose a stamp").pack(anchor="w", pady=(0, 6))
        self.listbox = tk.Listbox(body, selectmode="browse", activestyle="none", exportselection=False, height=14)
        for name in STAMP_NAMES:
            self.listbox.insert("end", name)
        self.listbox.selection_set(app.stamp_index)
        self.listbox.pack(fill="x")
        ttk.Label(body, text="Or type your own stamp text").pack(anchor="w", pady=(12, 4))
        self.custom = tk.StringVar(value=app.stamp_text)
        ttk.Entry(body, textvariable=self.custom, width=34).pack(anchor="w")
        bottom = ttk.Frame(self.win)
        bottom.pack(fill="x", padx=14, pady=(0, 14))
        ttk.Button(bottom, text="Use Stamp", command=self.use).pack(side="right")
        ttk.Button(bottom, text="Cancel", command=self.win.destroy).pack(side="right", padx=6)

    def use(self):
        chosen = self.listbox.curselection()
        if len(chosen) > 0:
            self.app.stamp_index = chosen[0]
        self.app.stamp_text = self.custom.get().strip()
        self.win.destroy()
        self.app.set_tool("stamp")


class App:
    def __init__(self):
        self.doc = None
        self.path = None
        self.password = None
        self.modified = False
        self.undo_stack = []
        self.redo_stack = []
        self.push_counter = 0
        self.recent = load_recent()
        self.color = (0.85, 0.1, 0.1)
        self.hl_color = (1.0, 0.9, 0.1)
        self.selection = None
        self.find_matches = []
        self.find_position = -1
        self.sign_mode = "text"
        self.sign_text = ""
        self.sign_path = SIGNATURE_FILE
        self.stamp_index = 0
        self.stamp_text = ""
        self.pending_image = None
        self.fullscreen = False
        self.panel_visible = True
        self.last_page = -1
        self.comment_map = {}
        self.root = tk.Tk()
        self.root.title(APP_NAME)
        center_window(self.root, 1220, 820)
        self.root.minsize(820, 520)
        self.tool_var = tk.StringVar(value="hand")
        self.line_width = tk.IntVar(value=2)
        self.font_size = tk.IntVar(value=14)
        self.page_var = tk.StringVar(value="1")
        self.total_var = tk.StringVar(value="/ 0")
        self.percent_var = tk.StringVar(value="100%")
        self.find_var = tk.StringVar(value="")
        self.find_info = tk.StringVar(value="")
        self.status_var = tk.StringVar(value="Open a PDF to begin")
        self.info_var = tk.StringVar(value="")
        self.build_menu()
        self.build_toolbar()
        self.build_status()
        self.build_body()
        self.build_find_bar()
        self.bind_keys()
        self.viewer.layout()
        self.root.protocol("WM_DELETE_WINDOW", self.exit_app)

    def add_item(self, menu, label, command, accel=""):
        menu.add_command(label=label, command=command, accelerator=accel)

    def build_menu(self):
        menubar = tk.Menu(self.root)
        self.root.config(menu=menubar)

        file_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="File", menu=file_menu)
        self.add_item(file_menu, "Open...", self.open_file, "Ctrl+O")
        self.add_item(file_menu, "Open Recent...", self.open_recent_window)
        file_menu.add_separator()
        create_menu = tk.Menu(file_menu, tearoff=0)
        file_menu.add_cascade(label="Create", menu=create_menu)
        self.add_item(create_menu, "Blank Document", self.new_blank)
        self.add_item(create_menu, "From Images...", self.create_from_images)
        self.add_item(create_menu, "From Text File...", self.create_from_text)
        self.add_item(create_menu, "Merge Files...", self.merge_files)
        file_menu.add_separator()
        self.add_item(file_menu, "Save", self.save_file, "Ctrl+S")
        self.add_item(file_menu, "Save As...", self.save_as, "Ctrl+Shift+S")
        export_menu = tk.Menu(file_menu, tearoff=0)
        file_menu.add_cascade(label="Export", menu=export_menu)
        self.add_item(export_menu, "Pages as Images...", self.export_images)
        self.add_item(export_menu, "Text File...", self.export_text)
        self.add_item(export_menu, "HTML File...", self.export_html)
        self.add_item(export_menu, "Word Document...", self.export_word)
        self.add_item(export_menu, "Pictures Inside the PDF...", self.extract_images)
        file_menu.add_separator()
        self.add_item(file_menu, "Protect with Password...", self.protect_document)
        self.add_item(file_menu, "Remove Password...", self.remove_password)
        self.add_item(file_menu, "Properties...", self.show_properties)
        self.add_item(file_menu, "Print...", self.print_document, "Ctrl+P")
        file_menu.add_separator()
        self.add_item(file_menu, "Close", self.close_file, "Ctrl+W")
        self.add_item(file_menu, "Exit", self.exit_app, "Ctrl+Q")

        edit_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Edit", menu=edit_menu)
        self.add_item(edit_menu, "Undo", self.undo, "Ctrl+Z")
        self.add_item(edit_menu, "Redo", self.redo, "Ctrl+Y")
        edit_menu.add_separator()
        self.add_item(edit_menu, "Copy Selected Text", self.copy_text, "Ctrl+C")
        self.add_item(edit_menu, "Select All Text on Page", self.select_page_text, "Ctrl+A")
        self.add_item(edit_menu, "Find...", self.show_find, "Ctrl+F")
        edit_menu.add_separator()
        self.add_item(edit_menu, "Edit Text", self.tool_edit)
        self.add_item(edit_menu, "Add Image...", self.tool_image)
        self.add_item(edit_menu, "Add Link", self.tool_link)

        view_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="View", menu=view_menu)
        self.add_item(view_menu, "Zoom In", self.zoom_in, "Ctrl++")
        self.add_item(view_menu, "Zoom Out", self.zoom_out, "Ctrl+-")
        self.add_item(view_menu, "Actual Size", self.zoom_actual, "Ctrl+1")
        self.add_item(view_menu, "Fit Width", self.zoom_fit_width, "Ctrl+2")
        self.add_item(view_menu, "Fit Page", self.zoom_fit_page, "Ctrl+0")
        view_menu.add_separator()
        self.add_item(view_menu, "Rotate All Pages Clockwise", self.rotate_view)
        view_menu.add_separator()
        self.add_item(view_menu, "First Page", self.first_page, "Home")
        self.add_item(view_menu, "Previous Page", self.prev_page, "PageUp")
        self.add_item(view_menu, "Next Page", self.next_page, "PageDown")
        self.add_item(view_menu, "Last Page", self.last_page_go, "End")
        view_menu.add_separator()
        self.add_item(view_menu, "Show or Hide Side Panel", self.toggle_panel)
        self.add_item(view_menu, "Full Screen", self.toggle_fullscreen, "F11")

        comment_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Comment", menu=comment_menu)
        self.add_item(comment_menu, "Highlight Text", self.tool_highlight)
        self.add_item(comment_menu, "Underline Text", self.tool_underline)
        self.add_item(comment_menu, "Strikethrough Text", self.tool_strike)
        self.add_item(comment_menu, "Squiggly Underline", self.tool_squiggly)
        comment_menu.add_separator()
        self.add_item(comment_menu, "Sticky Note", self.tool_note)
        self.add_item(comment_menu, "Text Box", self.tool_text)
        self.add_item(comment_menu, "Draw Freehand", self.tool_pen)
        self.add_item(comment_menu, "Rectangle", self.tool_rect)
        self.add_item(comment_menu, "Oval", self.tool_oval)
        self.add_item(comment_menu, "Line", self.tool_line)
        self.add_item(comment_menu, "Arrow", self.tool_arrow)
        self.add_item(comment_menu, "Stamp...", self.choose_stamp)
        comment_menu.add_separator()
        self.add_item(comment_menu, "Drawing Color...", self.choose_color)
        self.add_item(comment_menu, "Highlight Color...", self.choose_highlight_color)
        comment_menu.add_separator()
        self.add_item(comment_menu, "Delete All Comments", self.clear_comments)

        tools_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Tools", menu=tools_menu)
        self.add_item(tools_menu, "Organize Pages...", self.organize_pages)
        self.add_item(tools_menu, "Insert Pages From a File...", self.insert_pages)
        self.add_item(tools_menu, "Extract Pages...", self.extract_pages)
        self.add_item(tools_menu, "Delete Pages...", self.delete_pages)
        self.add_item(tools_menu, "Rotate Pages...", self.rotate_pages)
        self.add_item(tools_menu, "Split Document...", self.split_document)
        self.add_item(tools_menu, "Crop Pages", self.tool_crop)
        tools_menu.add_separator()
        self.add_item(tools_menu, "Header and Footer...", self.add_header_footer)
        self.add_item(tools_menu, "Watermark...", self.add_watermark)
        self.add_item(tools_menu, "Background Color...", self.set_background)
        self.add_item(tools_menu, "Bookmarks...", self.edit_bookmarks)
        tools_menu.add_separator()
        self.add_item(tools_menu, "Fill Form...", self.fill_form)
        self.add_item(tools_menu, "Add Text Field", self.tool_field_text)
        self.add_item(tools_menu, "Add Check Box", self.tool_field_check)
        self.add_item(tools_menu, "Sign...", self.start_sign)
        tools_menu.add_separator()
        self.add_item(tools_menu, "Mark for Redaction", self.tool_redact)
        self.add_item(tools_menu, "Apply Redactions", self.apply_redactions)
        tools_menu.add_separator()
        self.add_item(tools_menu, "Compress...", self.compress_document)
        self.add_item(tools_menu, "Recognize Text (OCR)...", self.run_ocr)
        self.add_item(tools_menu, "Compare With Another File...", self.compare_files)
        self.add_item(tools_menu, "Flatten Comments and Forms", self.flatten_document)
        self.add_item(tools_menu, "Remove Hidden Information", self.sanitize_document)

        help_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Help", menu=help_menu)
        self.add_item(help_menu, "Keyboard Shortcuts", self.show_shortcuts)
        self.add_item(help_menu, "About", self.show_about)

    def add_separator(self):
        ttk.Separator(self.toolbar, orient="vertical").pack(side="left", fill="y", padx=5, pady=3)

    def build_toolbar(self):
        self.toolbar = ttk.Frame(self.root)
        self.toolbar.pack(side="top", fill="x")
        ttk.Button(self.toolbar, text="Open", width=6, command=self.open_file).pack(side="left", padx=2, pady=3)
        ttk.Button(self.toolbar, text="Save", width=6, command=self.save_file).pack(side="left", padx=2, pady=3)
        ttk.Button(self.toolbar, text="Print", width=6, command=self.print_document).pack(side="left", padx=2, pady=3)
        self.add_separator()
        ttk.Button(self.toolbar, text="<", width=3, command=self.prev_page).pack(side="left", padx=1)
        self.page_entry = ttk.Entry(self.toolbar, textvariable=self.page_var, width=5, justify="center")
        self.page_entry.pack(side="left", padx=2)
        self.page_entry.bind("<Return>", self.goto_entry)
        ttk.Label(self.toolbar, textvariable=self.total_var, width=6).pack(side="left")
        ttk.Button(self.toolbar, text=">", width=3, command=self.next_page).pack(side="left", padx=1)
        self.add_separator()
        ttk.Button(self.toolbar, text="-", width=3, command=self.zoom_out).pack(side="left", padx=1)
        self.zoom_box = ttk.Combobox(self.toolbar, textvariable=self.percent_var, values=["25%", "50%", "75%", "100%", "125%", "150%", "200%", "300%", "400%"], width=6)
        self.zoom_box.pack(side="left", padx=2)
        self.zoom_box.bind("<<ComboboxSelected>>", self.zoom_from_box)
        self.zoom_box.bind("<Return>", self.zoom_from_box)
        ttk.Button(self.toolbar, text="+", width=3, command=self.zoom_in).pack(side="left", padx=1)
        ttk.Button(self.toolbar, text="Fit Width", command=self.zoom_fit_width).pack(side="left", padx=2)
        self.add_separator()
        tools = [
            ("Hand", "hand"),
            ("Select", "select"),
            ("Edit Text", "edit"),
            ("Highlight", "highlight"),
            ("Note", "note"),
            ("Text Box", "text"),
            ("Draw", "pen"),
            ("Redact", "redact"),
        ]
        for label, name in tools:
            tk.Radiobutton(self.toolbar, text=label, variable=self.tool_var, value=name, indicatoron=0, padx=7, pady=2, selectcolor="#bcd6f7", command=self.on_tool_change).pack(side="left", padx=1)
        ttk.Button(self.toolbar, text="Stamp", command=self.choose_stamp).pack(side="left", padx=2)
        ttk.Button(self.toolbar, text="Sign", command=self.start_sign).pack(side="left", padx=2)
        self.add_separator()
        self.color_swatch = tk.Label(self.toolbar, text="    ", bg=rgb_to_hex(self.color), relief="raised", cursor="hand2")
        self.color_swatch.pack(side="left", padx=2)
        self.color_swatch.bind("<Button-1>", self.choose_color)
        ttk.Label(self.toolbar, text="Width").pack(side="left", padx=(6, 2))
        ttk.Spinbox(self.toolbar, from_=1, to=20, width=3, textvariable=self.line_width).pack(side="left")
        ttk.Label(self.toolbar, text="Font").pack(side="left", padx=(6, 2))
        ttk.Spinbox(self.toolbar, from_=6, to=72, width=3, textvariable=self.font_size).pack(side="left")
        ttk.Button(self.toolbar, text="Find", width=6, command=self.show_find).pack(side="right", padx=4)

    def build_status(self):
        self.status_bar = ttk.Frame(self.root)
        self.status_bar.pack(side="bottom", fill="x")
        ttk.Label(self.status_bar, textvariable=self.status_var).pack(side="left", padx=8, pady=2)
        ttk.Label(self.status_bar, textvariable=self.info_var).pack(side="right", padx=8, pady=2)

    def build_body(self):
        self.panes = ttk.PanedWindow(self.root, orient="horizontal")
        self.panes.pack(fill="both", expand=True)
        self.left_frame = ttk.Frame(self.panes, width=210)
        self.panes.add(self.left_frame, weight=0)
        self.side_tabs = ttk.Notebook(self.left_frame)
        self.side_tabs.pack(fill="both", expand=True)
        self.thumbs = Thumbnails(self, self.side_tabs)
        self.side_tabs.add(self.thumbs.frame, text="Pages")

        bookmark_frame = ttk.Frame(self.side_tabs)
        self.side_tabs.add(bookmark_frame, text="Bookmarks")
        self.bm_tree = ttk.Treeview(bookmark_frame, show="tree")
        bm_bar = ttk.Scrollbar(bookmark_frame, orient="vertical", command=self.bm_tree.yview)
        self.bm_tree.configure(yscrollcommand=bm_bar.set)
        bm_bar.pack(side="right", fill="y")
        self.bm_tree.pack(side="left", fill="both", expand=True)
        self.bm_tree.bind("<<TreeviewSelect>>", self.on_bookmark_select)

        comment_frame = ttk.Frame(self.side_tabs)
        self.side_tabs.add(comment_frame, text="Comments")
        comment_buttons = ttk.Frame(comment_frame)
        comment_buttons.pack(side="bottom", fill="x", pady=4)
        ttk.Button(comment_buttons, text="Edit", command=self.comment_edit).pack(side="left", padx=4)
        ttk.Button(comment_buttons, text="Delete", command=self.comment_delete).pack(side="left", padx=4)
        self.cm_tree = ttk.Treeview(comment_frame, columns=("page", "kind", "text"), show="headings", selectmode="extended")
        self.cm_tree.heading("page", text="Pg")
        self.cm_tree.heading("kind", text="Type")
        self.cm_tree.heading("text", text="Comment")
        self.cm_tree.column("page", width=32, anchor="center")
        self.cm_tree.column("kind", width=70)
        self.cm_tree.column("text", width=100)
        cm_bar = ttk.Scrollbar(comment_frame, orient="vertical", command=self.cm_tree.yview)
        self.cm_tree.configure(yscrollcommand=cm_bar.set)
        cm_bar.pack(side="right", fill="y")
        self.cm_tree.pack(side="left", fill="both", expand=True)
        self.cm_tree.bind("<<TreeviewSelect>>", self.on_comment_select)

        self.viewer = Viewer(self, self.panes)
        self.panes.add(self.viewer.frame, weight=1)

    def build_find_bar(self):
        self.find_frame = ttk.Frame(self.root)
        ttk.Label(self.find_frame, text="Find:").pack(side="left", padx=(8, 4), pady=3)
        self.find_entry = ttk.Entry(self.find_frame, textvariable=self.find_var, width=32)
        self.find_entry.pack(side="left", padx=2)
        self.find_entry.bind("<Return>", self.run_find)
        ttk.Button(self.find_frame, text="Search", command=self.run_find).pack(side="left", padx=2)
        ttk.Button(self.find_frame, text="Previous", command=self.find_previous).pack(side="left", padx=2)
        ttk.Button(self.find_frame, text="Next", command=self.find_next).pack(side="left", padx=2)
        ttk.Label(self.find_frame, textvariable=self.find_info).pack(side="left", padx=10)
        ttk.Button(self.find_frame, text="Close", command=self.hide_find).pack(side="right", padx=8)

    def bind_to_keys(self, widget, key, function):
        mods = ["Control"]
        if sys.platform == "darwin":
            mods.append("Command")
        for mod in mods:
            widget.bind(f"<{mod}-{key}>", function)

    def bind_keys(self):
        self.bind_to_keys(self.root, "o", self.open_file)
        self.bind_to_keys(self.root, "s", self.save_file)
        self.bind_to_keys(self.root, "S", self.save_as)
        self.bind_to_keys(self.root, "w", self.close_file)
        self.bind_to_keys(self.root, "q", self.exit_app)
        self.bind_to_keys(self.root, "f", self.show_find)
        self.bind_to_keys(self.root, "p", self.print_document)
        self.bind_to_keys(self.root, "plus", self.zoom_in)
        self.bind_to_keys(self.root, "equal", self.zoom_in)
        self.bind_to_keys(self.root, "minus", self.zoom_out)
        self.bind_to_keys(self.root, "1", self.zoom_actual)
        self.bind_to_keys(self.root, "2", self.zoom_fit_width)
        self.bind_to_keys(self.root, "0", self.zoom_fit_page)
        self.root.bind("<F11>", self.toggle_fullscreen)
        self.root.bind("<Escape>", self.escape_key)
        canvas = self.viewer.canvas
        self.bind_to_keys(canvas, "z", self.undo)
        self.bind_to_keys(canvas, "y", self.redo)
        self.bind_to_keys(canvas, "c", self.copy_text)
        self.bind_to_keys(canvas, "a", self.select_page_text)
        canvas.bind("<Prior>", self.prev_page)
        canvas.bind("<Next>", self.next_page)
        canvas.bind("<Home>", self.first_page)
        canvas.bind("<End>", self.last_page_go)
        canvas.bind("<Up>", self.key_up)
        canvas.bind("<Down>", self.key_down)
        canvas.bind("<Left>", self.key_left)
        canvas.bind("<Right>", self.key_right)

    def key_up(self, event=None):
        self.viewer.canvas.yview_scroll(-3, "units")

    def key_down(self, event=None):
        self.viewer.canvas.yview_scroll(3, "units")

    def key_left(self, event=None):
        self.viewer.canvas.xview_scroll(-3, "units")

    def key_right(self, event=None):
        self.viewer.canvas.xview_scroll(3, "units")

    def escape_key(self, event=None):
        self.set_tool("hand")
        if self.fullscreen:
            self.toggle_fullscreen()

    def run(self):
        self.root.mainloop()

    def set_status(self, text):
        self.status_var.set(text)

    def need_doc(self):
        if self.doc is None:
            messagebox.showinfo(APP_NAME, "Open a PDF first.")
            return False
        return True

    def update_title(self):
        if self.doc is None:
            self.root.title(APP_NAME)
            return
        name = "Untitled"
        if self.path is not None:
            name = os.path.basename(self.path)
        mark = ""
        if self.modified:
            mark = "* "
        self.root.title(f"{mark}{name} - {APP_NAME}")

    def reload_all(self, keep=True):
        self.find_matches = []
        self.find_position = -1
        self.find_info.set("")
        self.viewer.find_marks = {}
        self.viewer.find_current = None
        self.viewer.sel_marks = {}
        self.selection = None
        self.viewer.reload(keep)
        self.thumbs.reload()
        self.update_panels()
        if self.doc is None:
            self.total_var.set("/ 0")
            self.page_var.set("1")
            self.info_var.set("")
        else:
            self.total_var.set(f"/ {self.doc.page_count}")
        self.last_page = -1
        self.page_changed()

    def page_changed(self):
        if self.doc is None:
            return
        index = self.viewer.current_page()
        if index != self.last_page:
            self.last_page = index
            self.thumbs.highlight(index)
            try:
                typing = self.root.focus_get() == self.page_entry
            except Exception:
                typing = False
            if not typing:
                self.page_var.set(str(index + 1))
        self.info_var.set(f"Page {index + 1} of {self.doc.page_count}     {self.viewer.percent}%")

    def confirm_discard(self):
        if self.doc is None or not self.modified:
            return True
        answer = messagebox.askyesnocancel("Unsaved changes", "Do you want to save your changes first?")
        if answer is None:
            return False
        if answer:
            return self.save_file()
        return True

    def add_recent(self, path):
        full = os.path.abspath(path)
        if full in self.recent:
            self.recent.remove(full)
        self.recent.insert(0, full)
        self.recent = self.recent[:10]
        save_recent(self.recent)

    def open_file(self, event=None):
        if not self.confirm_discard():
            return
        path = filedialog.askopenfilename(title="Open a PDF", filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")])
        if path == "" or len(path) == 0:
            return
        self.open_path(path)

    def open_path(self, path):
        try:
            doc = read_pdf(path)
        except Exception as e:
            messagebox.showerror("Open", f"This file could not be opened:\n{e}")
            return False
        password = None
        if doc.needs_pass:
            attempts = 0
            unlocked = False
            while attempts < 3:
                password = simpledialog.askstring("Password", "This file is password protected.\nEnter the password:", show="*", parent=self.root)
                if password is None:
                    break
                if doc.authenticate(password) > 0:
                    unlocked = True
                    break
                attempts = attempts + 1
                messagebox.showerror("Password", "That password is not correct.")
            if not unlocked:
                doc.close()
                return False
        if doc.page_count == 0:
            messagebox.showerror("Open", "This PDF has no pages.")
            doc.close()
            return False
        if self.doc is not None:
            self.doc.close()
        self.doc = doc
        self.path = path
        self.password = password
        self.undo_stack = []
        self.redo_stack = []
        self.modified = False
        self.add_recent(path)
        self.reload_all(False)
        self.update_title()
        self.set_status("Opened " + path)
        return True

    def set_new_document(self, new_doc):
        if self.doc is not None:
            self.doc.close()
        self.doc = new_doc
        self.path = None
        self.password = None
        self.undo_stack = []
        self.redo_stack = []
        self.modified = True
        self.reload_all(False)
        self.update_title()

    def open_recent_window(self):
        if len(self.recent) == 0:
            messagebox.showinfo("Open Recent", "No recent files yet.")
            return
        self.recent_win = make_window(self, "Open Recent", 660, 340)
        self.recent_list = tk.Listbox(self.recent_win, activestyle="none")
        self.recent_list.pack(fill="both", expand=True, padx=10, pady=10)
        for path in self.recent:
            self.recent_list.insert("end", path)
        self.recent_list.selection_set(0)
        self.recent_list.bind("<Double-Button-1>", self.recent_open)
        self.recent_list.bind("<Return>", self.recent_open)
        ttk.Button(self.recent_win, text="Open", command=self.recent_open).pack(pady=(0, 10))

    def recent_open(self, event=None):
        chosen = self.recent_list.curselection()
        if len(chosen) == 0:
            return
        path = self.recent[chosen[0]]
        self.recent_win.destroy()
        if not os.path.exists(path):
            messagebox.showerror("Open Recent", "That file no longer exists.")
            return
        if self.confirm_discard():
            self.open_path(path)

    def new_blank(self):
        if not self.confirm_discard():
            return
        new_doc = fitz.open()
        new_doc.new_page(width=595, height=842)
        self.set_new_document(new_doc)
        self.set_status("New blank document")

    def create_from_images(self):
        paths = filedialog.askopenfilenames(title="Choose images", filetypes=[("Images", "*.png *.jpg *.jpeg *.bmp *.gif *.tif *.tiff"), ("All files", "*.*")])
        if len(paths) == 0:
            return
        if not self.confirm_discard():
            return
        new_doc = fitz.open()
        progress = ProgressWindow(self, "Creating PDF", len(paths))
        try:
            number = 0
            for path in paths:
                image_doc = fitz.open(path)
                data = image_doc.convert_to_pdf()
                image_pdf = fitz.open(stream=data, filetype="pdf")
                new_doc.insert_pdf(image_pdf)
                image_pdf.close()
                image_doc.close()
                number = number + 1
                progress.step(number, f"Added {number} of {len(paths)}")
        except Exception as e:
            progress.close()
            messagebox.showerror("Create PDF", f"Could not use one of the images:\n{e}")
            return
        progress.close()
        self.set_new_document(new_doc)
        self.set_status("PDF created from images")

    def create_from_text(self):
        if not HAS_REPORTLAB:
            messagebox.showinfo("Create PDF", "This needs the reportlab library.\nInstall it with: pip install reportlab")
            return
        path = filedialog.askopenfilename(title="Choose a text file", filetypes=[("Text files", "*.txt"), ("All files", "*.*")])
        if path == "" or len(path) == 0:
            return
        if not self.confirm_discard():
            return
        try:
            file = open(path, "r", encoding="utf-8", errors="replace")
            text = file.read()
            file.close()
            text_to_pdf(text, PRINT_FILE)
            new_doc = read_pdf(PRINT_FILE)
        except Exception as e:
            messagebox.showerror("Create PDF", f"Could not create the PDF:\n{e}")
            return
        self.set_new_document(new_doc)
        self.set_status("PDF created from text file")

    def merge_files(self):
        MergeWindow(self)

    def save_file(self, event=None):
        if self.doc is None:
            return False
        if self.path is None:
            return self.save_as()
        return self.write_to(self.path)

    def save_as(self, event=None):
        if self.doc is None:
            return False
        name = "document.pdf"
        if self.path is not None:
            name = os.path.basename(self.path)
        path = filedialog.asksaveasfilename(title="Save As", defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")], initialfile=name)
        if path == "" or len(path) == 0:
            return False
        return self.write_to(path)

    def write_to(self, path):
        try:
            self.doc.save(path, garbage=3, deflate=True)
        except Exception as e:
            messagebox.showerror("Save", f"The file could not be saved:\n{e}")
            return False
        self.path = path
        self.modified = False
        self.add_recent(path)
        self.update_title()
        self.set_status("Saved " + path)
        return True

    def close_file(self, event=None):
        if self.doc is None:
            return
        if not self.confirm_discard():
            return
        self.doc.close()
        self.doc = None
        self.path = None
        self.password = None
        self.undo_stack = []
        self.redo_stack = []
        self.modified = False
        self.reload_all(False)
        self.update_title()
        self.set_status("Open a PDF to begin")

    def exit_app(self, event=None):
        if not self.confirm_discard():
            return
        self.root.destroy()

    def print_document(self, event=None):
        if not self.need_doc():
            return
        try:
            self.doc.save(PRINT_FILE)
            if sys.platform.startswith("win"):
                os.startfile(PRINT_FILE, "print")
            else:
                subprocess.run(["lp", PRINT_FILE])
        except Exception as e:
            messagebox.showerror("Print", f"Could not print:\n{e}")
            return
        self.set_status("Sent to the printer")

    def show_properties(self):
        if self.need_doc():
            PropertiesWindow(self)

    def protect_document(self):
        if self.need_doc():
            ProtectWindow(self)

    def remove_password(self):
        if not self.need_doc():
            return
        if not self.doc.is_encrypted:
            messagebox.showinfo("Remove Password", "This document is not protected.")
            return
        output = filedialog.asksaveasfilename(title="Save without password", defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")], initialfile="unprotected.pdf")
        if output == "" or len(output) == 0:
            return
        try:
            copy = fitz.open(stream=self.doc.tobytes(), filetype="pdf")
            if copy.needs_pass:
                copy.authenticate(self.password)
            copy.save(output, encryption=fitz.PDF_ENCRYPT_NONE)
            copy.close()
        except Exception as e:
            messagebox.showerror("Remove Password", f"Could not remove the password:\n{e}")
            return
        messagebox.showinfo("Remove Password", "A copy without a password was saved.")

    def base_name(self):
        if self.path is None:
            return "document"
        return os.path.splitext(os.path.basename(self.path))[0]

    def export_images(self):
        if not self.need_doc():
            return
        dpi = simpledialog.askinteger("Export Images", "Resolution in DPI (72 to 600):", initialvalue=150, minvalue=72, maxvalue=600, parent=self.root)
        if dpi is None:
            return
        folder = filedialog.askdirectory(title="Choose a folder for the images")
        if folder == "" or len(folder) == 0:
            return
        total = self.doc.page_count
        progress = ProgressWindow(self, "Exporting", total)
        try:
            for i in range(total):
                pix = self.doc[i].get_pixmap(dpi=dpi, alpha=False)
                pix.save(os.path.join(folder, f"{self.base_name()}_page{i + 1}.png"))
                progress.step(i + 1, f"Page {i + 1} of {total}")
        except Exception as e:
            progress.close()
            messagebox.showerror("Export", f"Could not export:\n{e}")
            return
        progress.close()
        messagebox.showinfo("Export", f"Saved {total} images in:\n{folder}")

    def export_text(self):
        if not self.need_doc():
            return
        path = filedialog.asksaveasfilename(title="Export Text", defaultextension=".txt", filetypes=[("Text files", "*.txt")], initialfile=self.base_name() + ".txt")
        if path == "" or len(path) == 0:
            return
        try:
            parts = []
            for i in range(self.doc.page_count):
                parts.append(f"--- Page {i + 1} ---")
                parts.append(self.doc[i].get_text("text"))
            file = open(path, "w", encoding="utf-8")
            file.write("\n".join(parts))
            file.close()
        except Exception as e:
            messagebox.showerror("Export", f"Could not export:\n{e}")
            return
        messagebox.showinfo("Export", "The text was saved.")

    def export_html(self):
        if not self.need_doc():
            return
        path = filedialog.asksaveasfilename(title="Export HTML", defaultextension=".html", filetypes=[("HTML files", "*.html")], initialfile=self.base_name() + ".html")
        if path == "" or len(path) == 0:
            return
        try:
            parts = ["<html><body>"]
            for i in range(self.doc.page_count):
                parts.append(self.doc[i].get_text("html"))
            parts.append("</body></html>")
            file = open(path, "w", encoding="utf-8")
            file.write("\n".join(parts))
            file.close()
        except Exception as e:
            messagebox.showerror("Export", f"Could not export:\n{e}")
            return
        messagebox.showinfo("Export", "The HTML file was saved.")

    def clean_text(self, text):
        result = ""
        for letter in text:
            if ord(letter) >= 32 or letter == "\n" or letter == "\t":
                result = result + letter
        return result

    def export_word(self):
        if not self.need_doc():
            return
        if not HAS_DOCX:
            messagebox.showinfo("Export", "This needs the python-docx library.\nInstall it with: pip install python-docx")
            return
        path = filedialog.asksaveasfilename(title="Export Word Document", defaultextension=".docx", filetypes=[("Word documents", "*.docx")], initialfile=self.base_name() + ".docx")
        if path == "" or len(path) == 0:
            return
        total = self.doc.page_count
        progress = ProgressWindow(self, "Exporting", total)
        try:
            document = docx.Document()
            for i in range(total):
                blocks = self.doc[i].get_text("blocks", sort=True)
                for block in blocks:
                    if block[6] != 0:
                        continue
                    text = self.clean_text(block[4]).strip()
                    if text != "":
                        document.add_paragraph(text)
                if i < total - 1:
                    document.add_page_break()
                progress.step(i + 1, f"Page {i + 1} of {total}")
            document.save(path)
        except Exception as e:
            progress.close()
            messagebox.showerror("Export", f"Could not export:\n{e}")
            return
        progress.close()
        messagebox.showinfo("Export", "The Word document was saved.\nOnly text is exported, not pictures or layout.")

    def extract_images(self):
        if not self.need_doc():
            return
        folder = filedialog.askdirectory(title="Choose a folder for the pictures")
        if folder == "" or len(folder) == 0:
            return
        count = 0
        try:
            for i in range(self.doc.page_count):
                images = self.doc[i].get_images(full=True)
                for k in range(len(images)):
                    info = self.doc.extract_image(images[k][0])
                    name = f"{self.base_name()}_page{i + 1}_image{k + 1}.{info['ext']}"
                    file = open(os.path.join(folder, name), "wb")
                    file.write(info["image"])
                    file.close()
                    count = count + 1
        except Exception as e:
            messagebox.showerror("Export", f"Could not extract the pictures:\n{e}")
            return
        messagebox.showinfo("Export", f"Saved {count} pictures.")

    def push_undo(self):
        data = self.doc.tobytes()
        self.undo_stack.append(data)
        if len(self.undo_stack) > MAX_UNDO:
            self.undo_stack.pop(0)
        self.redo_stack = []
        self.modified = True
        self.push_counter = self.push_counter + 1
        self.update_title()
        return self.push_counter

    def rollback(self, counter):
        if counter != self.push_counter:
            return
        if len(self.undo_stack) == 0:
            return
        data = self.undo_stack.pop()
        self.load_bytes(data)

    def load_bytes(self, data):
        new_doc = fitz.open(stream=data, filetype="pdf")
        if new_doc.needs_pass and self.password is not None:
            new_doc.authenticate(self.password)
        old = self.doc
        self.doc = new_doc
        if old is not None:
            old.close()
        self.reload_all(True)

    def undo(self, event=None):
        if self.doc is None:
            return
        if len(self.undo_stack) == 0:
            self.set_status("Nothing to undo")
            return
        self.redo_stack.append(self.doc.tobytes())
        data = self.undo_stack.pop()
        self.load_bytes(data)
        self.modified = True
        self.update_title()
        self.set_status("Undone")

    def redo(self, event=None):
        if self.doc is None:
            return
        if len(self.redo_stack) == 0:
            self.set_status("Nothing to redo")
            return
        self.undo_stack.append(self.doc.tobytes())
        data = self.redo_stack.pop()
        self.load_bytes(data)
        self.modified = True
        self.update_title()
        self.set_status("Redone")

    def copy_text(self, event=None):
        if self.selection is None:
            self.set_status("Nothing selected. Use the Select tool and drag over some text.")
            return
        self.root.clipboard_clear()
        self.root.clipboard_append(self.selection["text"])
        self.set_status("Text copied")

    def select_page_text(self, event=None):
        if self.doc is None:
            return
        index = self.viewer.current_page()
        page = self.doc[index]
        self.select_text(index, page, unrotated_rect(page))

    def zoom_in(self, event=None):
        if self.doc is None:
            return
        current = self.viewer.percent
        target = current * 1.25
        for level in ZOOM_LEVELS:
            if level > current:
                target = level
                break
        self.viewer.set_percent(int(target))

    def zoom_out(self, event=None):
        if self.doc is None:
            return
        current = self.viewer.percent
        target = current / 1.25
        for level in ZOOM_LEVELS[::-1]:
            if level < current:
                target = level
                break
        self.viewer.set_percent(int(target))

    def zoom_actual(self, event=None):
        self.viewer.set_percent(100)

    def zoom_fit_width(self, event=None):
        self.viewer.fit_width()

    def zoom_fit_page(self, event=None):
        self.viewer.fit_page()

    def zoom_from_box(self, event=None):
        text = self.percent_var.get().replace("%", "").strip()
        try:
            value = int(text)
        except ValueError:
            self.percent_var.set(f"{self.viewer.percent}%")
            return
        self.viewer.set_percent(value)

    def goto_entry(self, event=None):
        if self.doc is None:
            return
        try:
            number = int(self.page_var.get())
        except ValueError:
            self.page_var.set(str(self.viewer.current_page() + 1))
            return
        if number < 1:
            number = 1
        if number > self.doc.page_count:
            number = self.doc.page_count
        self.viewer.goto_page(number - 1)
        self.viewer.canvas.focus_set()

    def first_page(self, event=None):
        if self.doc is not None:
            self.viewer.goto_page(0)

    def last_page_go(self, event=None):
        if self.doc is not None:
            self.viewer.goto_page(self.doc.page_count - 1)

    def prev_page(self, event=None):
        if self.doc is not None:
            self.viewer.goto_page(self.viewer.current_page() - 1)

    def next_page(self, event=None):
        if self.doc is not None:
            self.viewer.goto_page(self.viewer.current_page() + 1)

    def rotate_view(self):
        if not self.need_doc():
            return
        counter = self.push_undo()
        try:
            for i in range(self.doc.page_count):
                page = self.doc[i]
                page.set_rotation((page.rotation + 90) % 360)
        except Exception as e:
            self.rollback(counter)
            messagebox.showerror("Rotate", str(e))
            return
        self.reload_all()

    def toggle_fullscreen(self, event=None):
        self.fullscreen = not self.fullscreen
        self.root.attributes("-fullscreen", self.fullscreen)

    def toggle_panel(self):
        if self.panel_visible:
            self.panes.forget(self.left_frame)
            self.panel_visible = False
        else:
            self.panes.insert(0, self.left_frame)
            self.panel_visible = True

    def show_find(self, event=None):
        if self.doc is None:
            return
        self.find_frame.pack(fill="x", before=self.panes)
        self.find_entry.focus_set()
        self.find_entry.select_range(0, "end")

    def hide_find(self):
        self.find_frame.pack_forget()
        self.find_matches = []
        self.find_position = -1
        self.find_info.set("")
        self.viewer.find_marks = {}
        self.viewer.find_current = None
        self.viewer.draw_overlays()

    def run_find(self, event=None):
        if self.doc is None:
            return
        needle = self.find_var.get()
        if needle == "":
            return
        self.find_matches = []
        self.viewer.find_marks = {}
        self.viewer.find_current = None
        for i in range(self.doc.page_count):
            page = self.doc[i]
            found = page.search_for(needle)
            if len(found) == 0:
                continue
            rects = []
            for rect in found:
                rects.append(rect * page.rotation_matrix)
            self.viewer.find_marks[i] = rects
            for n in range(len(rects)):
                self.find_matches.append((i, n))
        if len(self.find_matches) == 0:
            self.find_info.set("No matches")
            self.find_position = -1
            self.viewer.draw_overlays()
            return
        start = 0
        current = self.viewer.current_page()
        for k in range(len(self.find_matches)):
            if self.find_matches[k][0] >= current:
                start = k
                break
        self.go_to_match(start)

    def go_to_match(self, position):
        total = len(self.find_matches)
        if total == 0:
            return
        if position < 0:
            position = total - 1
        if position >= total:
            position = 0
        self.find_position = position
        index, n = self.find_matches[position]
        self.viewer.find_current = (index, n)
        self.viewer.draw_overlays()
        self.viewer.scroll_to_rect(index, self.viewer.find_marks[index][n])
        self.find_info.set(f"Match {position + 1} of {total}")

    def find_next(self):
        if len(self.find_matches) == 0:
            self.run_find()
            return
        self.go_to_match(self.find_position + 1)

    def find_previous(self):
        if len(self.find_matches) == 0:
            self.run_find()
            return
        self.go_to_match(self.find_position - 1)

    def update_panels(self):
        self.fill_bookmarks()
        self.fill_comments()

    def fill_bookmarks(self):
        for item in self.bm_tree.get_children():
            self.bm_tree.delete(item)
        if self.doc is None:
            return
        parents = {}
        try:
            toc = self.doc.get_toc()
        except Exception:
            toc = []
        for entry in toc:
            level = entry[0]
            parent = ""
            if level > 1:
                parent = parents.get(level - 1, "")
            item = self.bm_tree.insert(parent, "end", text=entry[1], values=(entry[2],), open=True)
            parents[level] = item

    def on_bookmark_select(self, event=None):
        chosen = self.bm_tree.selection()
        if len(chosen) == 0 or self.doc is None:
            return
        values = self.bm_tree.item(chosen[0], "values")
        try:
            number = int(values[0])
        except (ValueError, IndexError):
            return
        if number >= 1:
            self.viewer.goto_page(number - 1)

    def fill_comments(self):
        for item in self.cm_tree.get_children():
            self.cm_tree.delete(item)
        self.comment_map = {}
        if self.doc is None:
            return
        for i in range(self.doc.page_count):
            try:
                for annot in self.doc[i].annots():
                    content = annot.info.get("content", "")
                    if content is None:
                        content = ""
                    content = content.replace("\n", " ")
                    iid = self.cm_tree.insert("", "end", values=(i + 1, annot.type[1], content))
                    self.comment_map[iid] = (i, annot.xref)
            except Exception:
                continue

    def on_comment_select(self, event=None):
        chosen = self.cm_tree.selection()
        if len(chosen) == 0 or self.doc is None:
            return
        if chosen[0] in self.comment_map:
            self.viewer.goto_page(self.comment_map[chosen[0]][0])

    def find_annot(self, page_index, xref):
        page = self.doc[page_index]
        for annot in page.annots():
            if annot.xref == xref:
                return annot
        return None

    def comment_delete(self):
        if self.doc is None:
            return
        chosen = self.cm_tree.selection()
        if len(chosen) == 0:
            messagebox.showinfo("Comments", "Select a comment in the list first.")
            return
        counter = self.push_undo()
        try:
            for iid in chosen:
                if iid not in self.comment_map:
                    continue
                page_index, xref = self.comment_map[iid]
                annot = self.find_annot(page_index, xref)
                if annot is not None:
                    self.doc[page_index].delete_annot(annot)
        except Exception as e:
            self.rollback(counter)
            messagebox.showerror("Comments", str(e))
            return
        self.reload_all()
        self.set_status("Comment deleted")

    def comment_edit(self):
        if self.doc is None:
            return
        chosen = self.cm_tree.selection()
        if len(chosen) != 1 or chosen[0] not in self.comment_map:
            messagebox.showinfo("Comments", "Select one comment in the list first.")
            return
        page_index, xref = self.comment_map[chosen[0]]
        annot = self.find_annot(page_index, xref)
        if annot is None:
            return
        old = annot.info.get("content", "")
        text = simpledialog.askstring("Edit Comment", "Comment text:", initialvalue=old, parent=self.root)
        if text is None:
            return
        counter = self.push_undo()
        try:
            annot = self.find_annot(page_index, xref)
            annot.set_info(content=text)
            annot.update()
        except Exception as e:
            self.rollback(counter)
            messagebox.showerror("Comments", str(e))
            return
        self.reload_all()

    def clear_comments(self):
        if not self.need_doc():
            return
        if not messagebox.askyesno("Delete All Comments", "Remove every comment, highlight and drawing from this document?"):
            return
        counter = self.push_undo()
        try:
            for i in range(self.doc.page_count):
                page = self.doc[i]
                while True:
                    target = None
                    for annot in page.annots():
                        target = annot
                        break
                    if target is None:
                        break
                    page.delete_annot(target)
        except Exception as e:
            self.rollback(counter)
            messagebox.showerror("Comments", str(e))
            return
        self.reload_all()
        self.set_status("All comments deleted")

    def choose_color(self, event=None):
        result = colorchooser.askcolor(color=rgb_to_hex(self.color), parent=self.root, title="Choose a color")
        if result[0] is None:
            return
        self.color = (result[0][0] / 255, result[0][1] / 255, result[0][2] / 255)
        self.color_swatch.config(bg=rgb_to_hex(self.color))

    def choose_highlight_color(self):
        result = colorchooser.askcolor(color=rgb_to_hex(self.hl_color), parent=self.root, title="Highlight color")
        if result[0] is None:
            return
        self.hl_color = (result[0][0] / 255, result[0][1] / 255, result[0][2] / 255)

    def set_tool(self, name):
        self.tool_var.set(name)
        self.on_tool_change()

    def on_tool_change(self):
        tool = self.tool_var.get()
        self.viewer.set_cursor(tool)
        self.selection = None
        self.viewer.sel_marks = {}
        self.viewer.draw_overlays()
        self.set_status(TOOL_HINTS.get(tool, ""))

    def tool_hand(self):
        self.set_tool("hand")

    def tool_select(self):
        self.set_tool("select")

    def tool_edit(self):
        if self.need_doc():
            self.set_tool("edit")

    def tool_highlight(self):
        if self.need_doc():
            self.set_tool("highlight")

    def tool_underline(self):
        if self.need_doc():
            self.set_tool("underline")

    def tool_strike(self):
        if self.need_doc():
            self.set_tool("strike")

    def tool_squiggly(self):
        if self.need_doc():
            self.set_tool("squiggly")

    def tool_note(self):
        if self.need_doc():
            self.set_tool("note")

    def tool_text(self):
        if self.need_doc():
            self.set_tool("text")

    def tool_pen(self):
        if self.need_doc():
            self.set_tool("pen")

    def tool_rect(self):
        if self.need_doc():
            self.set_tool("rect")

    def tool_oval(self):
        if self.need_doc():
            self.set_tool("oval")

    def tool_line(self):
        if self.need_doc():
            self.set_tool("line")

    def tool_arrow(self):
        if self.need_doc():
            self.set_tool("arrow")

    def tool_link(self):
        if self.need_doc():
            self.set_tool("link")

    def tool_crop(self):
        if self.need_doc():
            self.set_tool("crop")

    def tool_redact(self):
        if self.need_doc():
            self.set_tool("redact")

    def tool_field_text(self):
        if self.need_doc():
            self.set_tool("field_text")

    def tool_field_check(self):
        if self.need_doc():
            self.set_tool("field_check")

    def tool_image(self):
        if not self.need_doc():
            return
        path = filedialog.askopenfilename(title="Choose an image", filetypes=[("Images", "*.png *.jpg *.jpeg *.bmp *.gif"), ("All files", "*.*")])
        if path == "" or len(path) == 0:
            return
        self.pending_image = path
        self.set_tool("image")

    def choose_stamp(self):
        if self.need_doc():
            StampWindow(self)

    def start_sign(self):
        if self.need_doc():
            SignWindow(self)

    def to_page_point(self, page, x, y):
        return fitz.Point(x, y) * page.derotation_matrix

    def finish_change(self, index):
        self.viewer.refresh_page(index)
        self.thumbs.refresh(index)
        self.fill_comments()

    def follow_link(self, index, vx, vy):
        page = self.doc[index]
        point = self.to_page_point(page, vx, vy)
        for link in page.get_links():
            if link["from"].contains(point):
                if link["kind"] == fitz.LINK_URI:
                    webbrowser.open(link["uri"])
                    return
                if link["kind"] == fitz.LINK_GOTO:
                    self.viewer.goto_page(link["page"])
                    return

    def handle_tool(self, index, x0, y0, x1, y1, points, clicked):
        if self.doc is None:
            return
        counter = self.push_counter
        try:
            self.run_tool(index, x0, y0, x1, y1, points, clicked)
        except Exception as e:
            if self.push_counter != counter:
                self.rollback(self.push_counter)
            messagebox.showerror("Error", f"That action did not work:\n{e}")

    def run_tool(self, index, x0, y0, x1, y1, points, clicked):
        tool = self.tool_var.get()
        page = self.doc[index]
        p0 = self.to_page_point(page, x0, y0)
        p1 = self.to_page_point(page, x1, y1)
        rect = make_rect(p0, p1)
        if tool == "select":
            if clicked:
                self.selection = None
                self.viewer.sel_marks = {}
                self.viewer.draw_overlays()
                self.set_status("Selection cleared")
            else:
                self.select_text(index, page, rect)
        elif tool == "highlight" or tool == "underline" or tool == "strike" or tool == "squiggly":
            if not clicked:
                self.add_markup(index, page, rect, tool)
        elif tool == "note":
            self.add_note(index, page, p0)
        elif tool == "text":
            self.add_text_box(index, page, p0, rect, clicked)
        elif tool == "pen":
            self.add_pen(index, page, points)
        elif tool == "rect" or tool == "oval":
            if not clicked:
                self.add_shape(index, page, rect, tool)
        elif tool == "line" or tool == "arrow":
            if not clicked:
                self.add_line(index, page, p0, p1, tool)
        elif tool == "stamp":
            self.add_stamp(index, page, p0, rect, clicked)
        elif tool == "sign":
            self.add_signature(index, page, p0, rect, clicked)
        elif tool == "edit":
            self.edit_text(index, page, p0, rect, clicked)
        elif tool == "redact":
            if not clicked:
                self.mark_redaction(index, page, rect)
        elif tool == "link":
            if not clicked:
                self.add_link(index, page, rect)
        elif tool == "crop":
            if not clicked:
                self.crop_pages(index, rect)
        elif tool == "image":
            self.add_picture(index, page, p0, rect, clicked)
        elif tool == "field_text" or tool == "field_check":
            if not clicked:
                self.add_form_field(index, page, rect, tool)

    def lines_in_rect(self, page, rect):
        lines = {}
        order = []
        for w in page.get_text("words"):
            box = fitz.Rect(w[0], w[1], w[2], w[3])
            if box.intersects(rect):
                key = (w[5], w[6])
                if key in lines:
                    lines[key] = union_rect(lines[key], box)
                else:
                    lines[key] = box
                    order.append(key)
        result = []
        for key in order:
            result.append(lines[key])
        return result

    def select_text(self, index, page, rect):
        text = ""
        rects = []
        last = None
        for w in page.get_text("words"):
            box = fitz.Rect(w[0], w[1], w[2], w[3])
            if box.intersects(rect):
                key = (w[5], w[6])
                if last is None:
                    text = w[4]
                elif key == last:
                    text = text + " " + w[4]
                else:
                    text = text + "\n" + w[4]
                last = key
                rects.append(box * page.rotation_matrix)
        if len(rects) == 0:
            self.selection = None
            self.viewer.sel_marks = {}
            self.viewer.draw_overlays()
            self.set_status("No text found there")
            return
        self.selection = {"page": index, "text": text}
        self.viewer.sel_marks = {index: rects}
        self.viewer.draw_overlays()
        self.set_status(f"{len(rects)} words selected. Press Ctrl+C to copy.")

    def add_markup(self, index, page, rect, tool):
        lines = self.lines_in_rect(page, rect)
        if len(lines) == 0:
            self.set_status("No text found there")
            return
        self.push_undo()
        if tool == "highlight":
            annot = page.add_highlight_annot(lines)
            annot.set_colors(stroke=self.hl_color)
        elif tool == "underline":
            annot = page.add_underline_annot(lines)
            annot.set_colors(stroke=self.color)
        elif tool == "strike":
            annot = page.add_strikeout_annot(lines)
            annot.set_colors(stroke=self.color)
        else:
            annot = page.add_squiggly_annot(lines)
            annot.set_colors(stroke=self.color)
        annot.update()
        self.finish_change(index)
        self.set_status("Comment added")

    def add_note(self, index, page, point):
        text = simpledialog.askstring("Sticky Note", "Type your note:", parent=self.root)
        if text is None or text == "":
            return
        self.push_undo()
        annot = page.add_text_annot(point, text)
        annot.set_colors(stroke=self.hl_color)
        annot.update()
        self.finish_change(index)
        self.set_status("Note added")

    def add_text_box(self, index, page, point, rect, clicked):
        text = simpledialog.askstring("Text Box", "Type the text:", parent=self.root)
        if text is None or text == "":
            return
        text = only_latin(text)
        size = self.font_size.get()
        if clicked:
            rows = text.split("\n")
            widest = 1
            for row in rows:
                widest = max(widest, len(row))
            width = widest * size * 0.6 + 14
            width = min(max(width, 80), 450)
            extra = int(widest * size * 0.6 / 450)
            height = size * 1.5 * (len(rows) + extra) + 12
            rect = fitz.Rect(point.x, point.y, point.x + width, point.y + height)
        self.push_undo()
        page.add_freetext_annot(rect, text, fontsize=size, fontname="helv", text_color=self.color)
        self.finish_change(index)
        self.set_status("Text box added")

    def add_pen(self, index, page, points):
        if len(points) < 2:
            return
        stroke = []
        for p in points:
            stroke.append(self.to_page_point(page, p[0], p[1]))
        self.push_undo()
        annot = page.add_ink_annot([stroke])
        annot.set_colors(stroke=self.color)
        annot.set_border(width=self.line_width.get())
        annot.update()
        self.finish_change(index)

    def add_shape(self, index, page, rect, tool):
        self.push_undo()
        if tool == "rect":
            annot = page.add_rect_annot(rect)
        else:
            annot = page.add_circle_annot(rect)
        annot.set_colors(stroke=self.color)
        annot.set_border(width=self.line_width.get())
        annot.update()
        self.finish_change(index)

    def add_line(self, index, page, p0, p1, tool):
        self.push_undo()
        annot = page.add_line_annot(p0, p1)
        if tool == "arrow":
            annot.set_line_ends(fitz.PDF_ANNOT_LE_NONE, fitz.PDF_ANNOT_LE_OPEN_ARROW)
        annot.set_colors(stroke=self.color)
        annot.set_border(width=self.line_width.get())
        annot.update()
        self.finish_change(index)

    def add_stamp(self, index, page, point, rect, clicked):
        if clicked:
            rect = fitz.Rect(point.x - 80, point.y - 30, point.x + 80, point.y + 30)
        self.push_undo()
        if self.stamp_text != "":
            page.add_freetext_annot(rect, only_latin(self.stamp_text), fontsize=18, fontname="helv", text_color=(0.8, 0, 0), align=1)
        else:
            page.add_stamp_annot(rect, stamp=self.stamp_index)
        self.finish_change(index)
        self.set_status("Stamp added")

    def add_signature(self, index, page, point, rect, clicked):
        if self.sign_mode == "image":
            if clicked:
                picture = Image.open(self.sign_path)
                ratio = picture.height / picture.width
                picture.close()
                width = 160
                height = width * ratio
                if height > 90:
                    height = 90
                    width = height / ratio
                rect = fitz.Rect(point.x - width / 2, point.y - height / 2, point.x + width / 2, point.y + height / 2)
            self.push_undo()
            page.insert_image(rect, filename=self.sign_path, keep_proportion=True, overlay=True)
        else:
            text = only_latin(self.sign_text)
            size = 26
            length = fitz.get_text_length(text, fontname="tiit", fontsize=size)
            self.push_undo()
            page.insert_text(fitz.Point(point.x - length / 2, point.y), text, fontsize=size, fontname="tiit", color=(0, 0, 0.55))
        self.finish_change(index)
        self.set_tool("hand")
        self.set_status("Signature placed")

    def has_redact_marks(self, page):
        for annot in page.annots():
            if annot.type[0] == fitz.PDF_ANNOT_REDACT:
                return True
        return False

    def edit_text(self, index, page, point, rect, clicked):
        if not clicked:
            text = simpledialog.askstring("Add Text", "Type the text:", parent=self.root)
            if text is None or text == "":
                return
            self.push_undo()
            result = page.insert_textbox(rect, only_latin(text), fontsize=self.font_size.get(), fontname="helv", color=self.color)
            if result < 0:
                self.rollback(self.push_counter)
                messagebox.showinfo("Add Text", "The text does not fit. Drag a bigger box.")
                return
            self.finish_change(index)
            return
        target = None
        info = page.get_text("dict")
        for block in info["blocks"]:
            if block["type"] != 0:
                continue
            for line in block["lines"]:
                if fitz.Rect(line["bbox"]).contains(point):
                    target = line
        if target is None:
            self.insert_new_text(index, page, point)
            return
        if self.has_redact_marks(page):
            messagebox.showinfo("Edit Text", "This page has areas marked for redaction.\nApply them (Tools > Apply Redactions) or delete them first.")
            return
        spans = target["spans"]
        old = ""
        for span in spans:
            old = old + span["text"]
        new = simpledialog.askstring("Edit Text", "Change the text:", initialvalue=old, parent=self.root)
        if new is None or new == old:
            return
        first = spans[0]
        font = pick_font(first["font"])
        size = first["size"]
        color = int_to_rgb(first["color"])
        origin = fitz.Point(first["origin"])
        box = fitz.Rect(target["bbox"])
        self.push_undo()
        page.add_redact_annot(box, fill=(1, 1, 1))
        page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)
        page.insert_text(origin, only_latin(new), fontsize=size, fontname=font, color=color)
        self.finish_change(index)
        self.set_status("Text changed")

    def insert_new_text(self, index, page, point):
        text = simpledialog.askstring("Add Text", "Type the text to add here:", parent=self.root)
        if text is None or text == "":
            return
        size = self.font_size.get()
        self.push_undo()
        page.insert_text(fitz.Point(point.x, point.y + size), only_latin(text), fontsize=size, fontname="helv", color=self.color)
        self.finish_change(index)
        self.set_status("Text added")

    def mark_redaction(self, index, page, rect):
        self.push_undo()
        page.add_redact_annot(rect, fill=(0, 0, 0))
        self.finish_change(index)
        self.set_status("Marked. Use Tools > Apply Redactions to remove the content for good.")

    def add_link(self, index, page, rect):
        target = simpledialog.askstring("Add Link", "Type a web address or a page number:", parent=self.root)
        if target is None or target.strip() == "":
            return
        target = target.strip()
        if target.isnumeric():
            number = int(target)
            if number < 1 or number > self.doc.page_count:
                messagebox.showerror("Add Link", "That page does not exist.")
                return
            link = {"kind": fitz.LINK_GOTO, "from": rect, "page": number - 1, "to": fitz.Point(0, 0), "zoom": 0}
        else:
            if "://" not in target and not target.startswith("mailto:"):
                target = "https://" + target
            link = {"kind": fitz.LINK_URI, "from": rect, "uri": target}
        self.push_undo()
        page.insert_link(link)
        self.finish_change(index)
        self.set_status("Link added. Use the Hand tool to try it.")

    def crop_pages(self, index, rect):
        answer = messagebox.askyesnocancel("Crop Pages", "Crop every page to this area?\nYes = all pages. No = only this page.")
        if answer is None:
            return
        targets = [index]
        if answer:
            targets = list(range(self.doc.page_count))
        self.push_undo()
        for i in targets:
            target = self.doc[i]
            box = fitz.Rect(rect)
            box.intersect(unrotated_rect(target))
            if box.is_empty:
                continue
            target.set_cropbox(box)
        self.reload_all()
        self.set_status("Pages cropped")

    def add_picture(self, index, page, point, rect, clicked):
        path = self.pending_image
        if path is None:
            return
        if clicked:
            picture = Image.open(path)
            w, h = picture.size
            picture.close()
            width = min(w * 0.75, 300)
            height = width * h / w
            rect = fitz.Rect(point.x, point.y, point.x + width, point.y + height)
        self.push_undo()
        page.insert_image(rect, filename=path, keep_proportion=True)
        self.finish_change(index)
        self.set_tool("hand")
        self.set_status("Image added")

    def add_form_field(self, index, page, rect, tool):
        count = 1
        for i in range(self.doc.page_count):
            for widget in self.doc[i].widgets():
                count = count + 1
        name = simpledialog.askstring("Form Field", "Name for this field:", initialvalue=f"field{count}", parent=self.root)
        if name is None or name.strip() == "":
            return
        self.push_undo()
        widget = fitz.Widget()
        widget.field_name = name.strip()
        widget.rect = rect
        if tool == "field_text":
            widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
            widget.field_value = ""
            widget.text_fontsize = 12
        else:
            widget.field_type = fitz.PDF_WIDGET_TYPE_CHECKBOX
            widget.field_value = False
        page.add_widget(widget)
        self.finish_change(index)
        self.set_status("Form field added")

    def organize_pages(self):
        if self.need_doc():
            OrganizeWindow(self)

    def split_document(self):
        if self.need_doc():
            SplitWindow(self)

    def add_header_footer(self):
        if self.need_doc():
            HeaderFooterWindow(self)

    def add_watermark(self):
        if self.need_doc():
            WatermarkWindow(self)

    def edit_bookmarks(self):
        if self.need_doc():
            BookmarkWindow(self)

    def ask_pages(self, title, prompt):
        total = self.doc.page_count
        text = simpledialog.askstring(title, f"{prompt}\nThe document has {total} pages. Example: 1-3, 5, 8-", parent=self.root)
        if text is None:
            return None
        try:
            pages = parse_pages(text, total)
        except ValueError:
            messagebox.showerror(title, "Please type page numbers like 1-3, 5, 8-")
            return None
        if len(pages) == 0:
            messagebox.showinfo(title, "No valid pages were chosen.")
            return None
        return pages

    def insert_pages(self):
        if not self.need_doc():
            return
        path = filedialog.askopenfilename(title="Choose the PDF to insert", filetypes=[("PDF files", "*.pdf")])
        if path == "" or len(path) == 0:
            return
        try:
            source = read_pdf(path)
            if source.needs_pass:
                password = simpledialog.askstring("Password", "Password for that file:", show="*", parent=self.root)
                if password is None or source.authenticate(password) == 0:
                    messagebox.showerror("Insert Pages", "Could not open that file.")
                    return
        except Exception as e:
            messagebox.showerror("Insert Pages", f"Could not open that file:\n{e}")
            return
        count = source.page_count
        text = simpledialog.askstring("Insert Pages", f"Which pages of that file? It has {count}.\nLeave blank for all pages.", parent=self.root)
        if text is None:
            return
        try:
            pages = parse_pages(text, count)
        except ValueError:
            messagebox.showerror("Insert Pages", "Please type page numbers like 1-3, 5")
            return
        if len(pages) == 0:
            pages = list(range(count))
        total = self.doc.page_count
        position = simpledialog.askinteger("Insert Pages", f"Insert after which page? (0 = at the start, {total} = at the end)", initialvalue=self.viewer.current_page() + 1, minvalue=0, maxvalue=total, parent=self.root)
        if position is None:
            return
        counter = self.push_undo()
        try:
            for k in range(len(pages)):
                self.doc.insert_pdf(source, from_page=pages[k], to_page=pages[k], start_at=position + k)
        except Exception as e:
            self.rollback(counter)
            messagebox.showerror("Insert Pages", str(e))
            return
        source.close()
        self.reload_all()
        self.set_status(f"Inserted {len(pages)} pages")

    def extract_pages(self):
        if not self.need_doc():
            return
        pages = self.ask_pages("Extract Pages", "Which pages should be extracted?")
        if pages is None:
            return
        output = filedialog.asksaveasfilename(title="Save the extracted pages", defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")], initialfile="extracted.pdf")
        if output == "" or len(output) == 0:
            return
        try:
            new_doc = fitz.open()
            for index in pages:
                new_doc.insert_pdf(self.doc, from_page=index, to_page=index)
            new_doc.save(output)
            new_doc.close()
        except Exception as e:
            messagebox.showerror("Extract Pages", str(e))
            return
        messagebox.showinfo("Extract Pages", f"Saved {len(pages)} pages to a new file.")

    def delete_pages(self):
        if not self.need_doc():
            return
        pages = self.ask_pages("Delete Pages", "Which pages should be deleted?")
        if pages is None:
            return
        if len(pages) >= self.doc.page_count:
            messagebox.showinfo("Delete Pages", "A document needs at least one page.")
            return
        if not messagebox.askyesno("Delete Pages", f"Delete {len(pages)} page(s)?"):
            return
        pages.sort()
        counter = self.push_undo()
        try:
            for index in pages[::-1]:
                self.doc.delete_page(index)
        except Exception as e:
            self.rollback(counter)
            messagebox.showerror("Delete Pages", str(e))
            return
        self.reload_all()
        self.set_status(f"Deleted {len(pages)} pages")

    def rotate_pages(self):
        if not self.need_doc():
            return
        pages = self.ask_pages("Rotate Pages", "Which pages should be rotated?")
        if pages is None:
            return
        angle = simpledialog.askinteger("Rotate Pages", "Degrees clockwise (90, 180, 270 or -90):", initialvalue=90, parent=self.root)
        if angle is None:
            return
        if angle % 90 != 0:
            messagebox.showerror("Rotate Pages", "The angle must be a multiple of 90.")
            return
        counter = self.push_undo()
        try:
            for index in pages:
                page = self.doc[index]
                page.set_rotation((page.rotation + angle) % 360)
        except Exception as e:
            self.rollback(counter)
            messagebox.showerror("Rotate Pages", str(e))
            return
        self.reload_all()

    def set_background(self):
        if not self.need_doc():
            return
        result = colorchooser.askcolor(parent=self.root, title="Background color")
        if result[0] is None:
            return
        fill = (result[0][0] / 255, result[0][1] / 255, result[0][2] / 255)
        counter = self.push_undo()
        try:
            for i in range(self.doc.page_count):
                page = self.doc[i]
                page.draw_rect(unrotated_rect(page), color=None, fill=fill, overlay=False)
        except Exception as e:
            self.rollback(counter)
            messagebox.showerror("Background", str(e))
            return
        self.reload_all()
        self.set_status("Background added")

    def fill_form(self):
        if not self.need_doc():
            return
        count = 0
        for i in range(self.doc.page_count):
            for widget in self.doc[i].widgets():
                count = count + 1
        if count == 0:
            messagebox.showinfo("Fill Form", "This document has no form fields.\nUse Tools > Add Text Field or Add Check Box to make some.")
            return
        FormWindow(self)

    def apply_redactions(self):
        if not self.need_doc():
            return
        count = 0
        for i in range(self.doc.page_count):
            for annot in self.doc[i].annots():
                if annot.type[0] == fitz.PDF_ANNOT_REDACT:
                    count = count + 1
        if count == 0:
            messagebox.showinfo("Redact", "Nothing is marked for redaction.\nUse Tools > Mark for Redaction first.")
            return
        if not messagebox.askyesno("Apply Redactions", f"Permanently remove the content under {count} marked area(s)?\nYou can still undo this until you save."):
            return
        counter = self.push_undo()
        try:
            for i in range(self.doc.page_count):
                self.doc[i].apply_redactions()
        except Exception as e:
            self.rollback(counter)
            messagebox.showerror("Redact", str(e))
            return
        self.reload_all()
        self.set_status("Redactions applied")

    def compress_document(self):
        if self.need_doc():
            CompressWindow(self)

    def run_ocr(self):
        if not self.need_doc():
            return
        if not HAS_OCR:
            messagebox.showinfo("OCR", "This needs the pytesseract library.\nInstall it with: pip install pytesseract\nYou also need the Tesseract program installed.")
            return
        try:
            pytesseract.get_tesseract_version()
        except Exception:
            messagebox.showinfo("OCR", "The Tesseract program was not found.\nInstall Tesseract OCR and make sure it is on your PATH.")
            return
        if not messagebox.askyesno("Recognize Text", "Add a searchable text layer to every page?\nThis can take a while."):
            return
        scale = 200 / 72
        total = self.doc.page_count
        progress = ProgressWindow(self, "Recognizing text", total)
        counter = self.push_undo()
        added = 0
        try:
            for i in range(total):
                page = self.doc[i]
                pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                picture = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                data = pytesseract.image_to_data(picture, output_type=pytesseract.Output.DICT)
                for n in range(len(data["text"])):
                    word = data["text"][n].strip()
                    if word == "":
                        continue
                    try:
                        confidence = float(data["conf"][n])
                    except ValueError:
                        continue
                    if confidence < 30:
                        continue
                    left = data["left"][n] / scale
                    top = data["top"][n] / scale
                    height = data["height"][n] / scale
                    size = max(height * 0.8, 4)
                    view_point = fitz.Point(left, top + height * 0.82)
                    point = view_point * page.derotation_matrix
                    page.insert_text(point, only_latin(word), fontsize=size, fontname="helv", render_mode=3)
                    added = added + 1
                progress.step(i + 1, f"Page {i + 1} of {total}")
        except Exception as e:
            progress.close()
            self.rollback(counter)
            messagebox.showerror("OCR", f"Text recognition failed:\n{e}")
            return
        progress.close()
        self.reload_all()
        messagebox.showinfo("OCR", f"Added {added} words. The text can now be searched and copied.")

    def compare_files(self):
        if not self.need_doc():
            return
        path = filedialog.askopenfilename(title="Choose the file to compare with", filetypes=[("PDF files", "*.pdf")])
        if path == "" or len(path) == 0:
            return
        CompareWindow(self, path)

    def flatten_document(self):
        if not self.need_doc():
            return
        if not messagebox.askyesno("Flatten", "Make comments, drawings and form fields a permanent part of the pages?"):
            return
        counter = self.push_undo()
        note = "Comments and form fields are now part of the pages."
        try:
            self.doc.bake(annots=True, widgets=True)
        except Exception:
            try:
                new_doc = fitz.open()
                for i in range(self.doc.page_count):
                    page = self.doc[i]
                    pix = page.get_pixmap(dpi=150, alpha=False)
                    new_page = new_doc.new_page(width=page.rect.width, height=page.rect.height)
                    new_page.insert_image(new_page.rect, pixmap=pix)
                self.load_bytes(new_doc.tobytes())
                new_doc.close()
                note = "Your PyMuPDF version cannot merge layers, so every page was turned into a picture."
            except Exception as e:
                self.rollback(counter)
                messagebox.showerror("Flatten", str(e))
                return
        self.reload_all()
        messagebox.showinfo("Flatten", note)

    def sanitize_document(self):
        if not self.need_doc():
            return
        if not messagebox.askyesno("Remove Hidden Information", "This removes the document properties, all comments, bookmarks and attached files.\nContinue?"):
            return
        counter = self.push_undo()
        try:
            self.doc.set_metadata({})
            self.doc.set_toc([])
            for i in range(self.doc.page_count):
                page = self.doc[i]
                while True:
                    target = None
                    for annot in page.annots():
                        target = annot
                        break
                    if target is None:
                        break
                    page.delete_annot(target)
            for name in self.doc.embfile_names():
                self.doc.embfile_del(name)
            try:
                self.doc.del_xml_metadata()
            except Exception:
                pass
        except Exception as e:
            self.rollback(counter)
            messagebox.showerror("Remove Hidden Information", str(e))
            return
        self.reload_all()
        self.set_status("Hidden information removed")

    def show_shortcuts(self):
        lines = [
            "Ctrl+O          Open a file",
            "Ctrl+S          Save",
            "Ctrl+Shift+S    Save As",
            "Ctrl+W          Close the file",
            "Ctrl+P          Print",
            "Ctrl+F          Find text",
            "Ctrl+Z / Ctrl+Y Undo / Redo",
            "Ctrl+C          Copy selected text",
            "Ctrl+A          Select all text on the page",
            "Ctrl + / -      Zoom in / out",
            "Ctrl+1 / 2 / 0  Actual size / Fit width / Fit page",
            "Ctrl + mouse wheel   Zoom",
            "Page Up / Down  Previous / next page",
            "Home / End      First / last page",
            "F11             Full screen",
            "Esc             Back to the Hand tool",
            "",
            "On a Mac, use the Command key instead of Ctrl.",
        ]
        messagebox.showinfo("Keyboard Shortcuts", "\n".join(lines))

    def show_about(self):
        text = f"{APP_NAME}\n\nA simple PDF reader and editor made with Python.\nBuilt with Tkinter, PyMuPDF, pypdf, Pillow,\nreportlab, python-docx and pytesseract."
        messagebox.showinfo("About", text)


app = App()
if len(sys.argv) > 1:
    if os.path.exists(sys.argv[1]):
        app.open_path(sys.argv[1])
app.run()
